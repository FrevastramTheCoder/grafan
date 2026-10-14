from fastapi import FastAPI, HTTPException, UploadFile, File, Form, BackgroundTasks, Depends, Header
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel
from typing import Optional, Dict, Any
from io import BytesIO
import base64
from datetime import datetime, timezone
from contextlib import asynccontextmanager
import os
import sys
import asyncio
import re
import threading
import uuid
import hmac
import hashlib
import json
from pathlib import Path
from urllib.parse import urlparse
from apscheduler.schedulers.asyncio import AsyncIOScheduler

from sqlalchemy import (
    create_engine,
    Column,
    Integer,
    String,
    Text,
    DateTime,
    func,
    ForeignKey,
)
from sqlalchemy.orm import declarative_base, sessionmaker

from scraper import JobScraper
from cv_parser import CVParser
from email_sender import EmailSender
from cv_ai import (
    CVAIError,
    PROFILE_FIELDS as CV_PROFILE_FIELDS,
    configured as cv_ai_configured,
    update_profile as update_cv_profile,
)

# Windows may run Python with a legacy console encoding that cannot print log emojis.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# ============================================================
# DATABASE SETUP
# ============================================================

BACKEND_DIR = Path(__file__).resolve().parent

DEFAULT_DATABASE_PATH = (
    BACKEND_DIR / "instance" / "grafan.db"
)

DEFAULT_DATABASE_PATH.parent.mkdir(
    parents=True,
    exist_ok=True
)

DATABASE_URL = os.environ.get("DATABASE_URL")

if not DATABASE_URL:
    DATABASE_URL = (
        f"sqlite:///{DEFAULT_DATABASE_PATH.as_posix()}"
    )

elif DATABASE_URL.startswith("sqlite:///"):
    configured_path = Path(
        DATABASE_URL.removeprefix("sqlite:///")
    )

    if not configured_path.is_absolute():
        DATABASE_URL = (
            f"sqlite:///{(BACKEND_DIR / configured_path).resolve().as_posix()}"
        )

engine_options = (
    {"connect_args": {"check_same_thread": False}}
    if DATABASE_URL.startswith("sqlite")
    else {}
)

# SSL kwa PostgreSQL (Neon/Supabase)
if DATABASE_URL.startswith("postgresql"):
    if "sslmode" not in DATABASE_URL:
        DATABASE_URL += "?sslmode=require"

engine = create_engine(
    DATABASE_URL,
    **engine_options,
    pool_pre_ping=True,
)

SessionLocal = sessionmaker(
    bind=engine,
    autoflush=False,
    autocommit=False
)

Base = declarative_base()

UPLOAD_FOLDER = str(
    BACKEND_DIR / "uploads"
)

os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024

scrape_lock = threading.Lock()
scrape_runs = {}


# ============================================================
# MODELS
# ============================================================

class Job(Base):
    __tablename__ = "job"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    title = Column(
        String(200),
        nullable=False
    )

    company = Column(
        String(100)
    )

    job_type = Column(
        String(50)
    )

    job_category = Column(
        String(100)
    )

    location = Column(
        String(100)
    )

    salary = Column(
        String(100)
    )

    description = Column(
        Text
    )

    source = Column(
        String(100)
    )

    source_url = Column(
        Text
    )

    application_url = Column(
        Text
    )

    scraped_at = Column(
        DateTime
    )

    requirements = Column(
        Text
    )

    responsibilities = Column(
        Text
    )

    qualifications = Column(
        Text
    )

    required_skills = Column(
        Text
    )

    experience_requirements = Column(
        Text
    )

    application_deadline = Column(
        String(100)
    )

    application_instructions = Column(
        Text
    )

    job_status = Column(
        String(30),
        default="active",
        nullable=False
    )

    completeness_score = Column(
        Integer,
        default=0
    )

    duplicate_of = Column(
        Integer
    )

    recruiter_email = Column(
        String(100)
    )

    posted_date = Column(
        DateTime,
        default=datetime.utcnow
    )


class Application(Base):
    __tablename__ = "application"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    job_id = Column(
        Integer
    )

    applicant_name = Column(
        String(100)
    )

    applicant_email = Column(
        String(100)
    )

    applicant_phone = Column(
        String(20)
    )

    cv_filename = Column(
        String(200)
    )

    cover_letter = Column(
        Text
    )

    applied_date = Column(
        DateTime,
        default=datetime.utcnow
    )


class SavedJob(Base):
    __tablename__ = "saved_job"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    owner_key = Column(
        String(64),
        nullable=False,
        index=True
    )

    job_id = Column(
        Integer,
        nullable=False,
        index=True
    )

    saved_at = Column(
        DateTime,
        default=datetime.utcnow
    )


class UserProfile(Base):
    __tablename__ = "user_profile"

    key_hash = Column(
        String(64),
        primary_key=True
    )

    profile_json = Column(
        Text,
        nullable=False
    )

    updated_at = Column(
        DateTime,
        default=datetime.utcnow,
        nullable=False
    )


class BusinessProfile(Base):
    __tablename__ = "business_profile"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    owner_hash = Column(
        String(64),
        nullable=False,
        unique=True,
        index=True
    )

    slug = Column(
        String(100),
        nullable=False,
        unique=True,
        index=True
    )

    profile_json = Column(
        Text,
        nullable=False
    )

    created_at = Column(
        DateTime,
        default=datetime.utcnow,
        nullable=False
    )

    updated_at = Column(
        DateTime,
        default=datetime.utcnow,
        nullable=False
    )


class Advertisement(Base):
    __tablename__ = "advertisement"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    owner_hash = Column(
        String(64),
        nullable=False,
        index=True
    )

    business_id = Column(
        Integer,
        ForeignKey("business_profile.id"),
        nullable=True,
        index=True
    )

    job_id = Column(
        Integer,
        ForeignKey("job.id"),
        nullable=True,
        index=True
    )

    title = Column(
        String(200),
        nullable=False
    )

    category = Column(
        String(50),
        nullable=False,
        index=True
    )

    description = Column(
        Text,
        nullable=False
    )

    status = Column(
        String(30),
        default="draft",
        nullable=False,
        index=True
    )

    rejection_reason = Column(
        Text
    )

    payload_json = Column(
        Text,
        nullable=False
    )

    promotion = Column(
        String(20),
        default="free",
        nullable=False
    )

    featured = Column(
        Integer,
        default=0,
        nullable=False
    )

    created_at = Column(
        DateTime,
        default=datetime.utcnow,
        nullable=False
    )

    updated_at = Column(
        DateTime,
        default=datetime.utcnow,
        nullable=False
    )

    published_at = Column(
        DateTime
    )

    expires_at = Column(
        DateTime
    )


class AdvertisementEvent(Base):
    __tablename__ = "advertisement_event"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    advertisement_id = Column(
        Integer,
        ForeignKey("advertisement.id"),
        nullable=False,
        index=True
    )

    event_type = Column(
        String(30),
        nullable=False,
        index=True
    )

    created_at = Column(
        DateTime,
        default=datetime.utcnow
    )


class AdvertisementImage(Base):
    __tablename__ = "advertisement_image"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    advertisement_id = Column(
        Integer,
        ForeignKey("advertisement.id"),
        nullable=False,
        index=True
    )

    kind = Column(
        String(20),
        nullable=False
    )

    image_data = Column(
        Text,
        nullable=False
    )

    sort_index = Column(
        Integer,
        default=0,
        nullable=False
    )


Base.metadata.create_all(
    bind=engine
)


# ============================================================
# ADDITIVE DATABASE MIGRATION
# ============================================================

def migrate_additive_job_columns():
    """
    Add nullable listing metadata while preserving every
    existing row.
    """

    if engine.dialect.name != "sqlite":
        return

    with engine.begin() as connection:

        columns = {
            row["name"]
            for row in connection.exec_driver_sql(
                "PRAGMA table_info(job)"
            ).mappings()
        }

        for name, sql_type in (
            ("source_url", "TEXT"),
            ("application_url", "TEXT"),
            ("scraped_at", "DATETIME"),
            ("requirements", "TEXT"),
            ("responsibilities", "TEXT"),
            ("qualifications", "TEXT"),
            ("required_skills", "TEXT"),
            ("experience_requirements", "TEXT"),
            ("application_deadline", "TEXT"),
            ("application_instructions", "TEXT"),
            (
                "job_status",
                "TEXT NOT NULL DEFAULT 'active'"
            ),
            (
                "completeness_score",
                "INTEGER NOT NULL DEFAULT 0"
            ),
            ("duplicate_of", "INTEGER"),
            ("job_category", "TEXT"),
        ):

            if name not in columns:

                connection.exec_driver_sql(
                    f"ALTER TABLE job ADD COLUMN {name} {sql_type}"
                )


migrate_additive_job_columns()


# ============================================================
# PYDANTIC SCHEMAS
# ============================================================

class JobCreate(BaseModel):
    title: str
    company: Optional[str] = None
    job_type: Optional[str] = None
    location: Optional[str] = None
    salary: Optional[str] = None
    description: Optional[str] = ""
    source: Optional[str] = "Manual Entry"
    source_url: Optional[str] = ""
    application_url: Optional[str] = ""
    recruiter_email: Optional[str] = ""
    requirements: Optional[str] = None
    responsibilities: Optional[str] = None
    qualifications: Optional[str] = None
    required_skills: Optional[str] = None
    experience_requirements: Optional[str] = None
    application_deadline: Optional[str] = None
    application_instructions: Optional[str] = None


class JobUpdate(BaseModel):
    title: Optional[str] = None
    company: Optional[str] = None
    job_type: Optional[str] = None
    location: Optional[str] = None
    salary: Optional[str] = None
    description: Optional[str] = None
    source_url: Optional[str] = None
    application_url: Optional[str] = None
    recruiter_email: Optional[str] = None
    requirements: Optional[str] = None
    responsibilities: Optional[str] = None
    qualifications: Optional[str] = None
    required_skills: Optional[str] = None
    experience_requirements: Optional[str] = None
    application_deadline: Optional[str] = None
    application_instructions: Optional[str] = None
    job_status: Optional[str] = None


class AIQuestion(BaseModel):
    question: str
    skills: Optional[str] = ""


class SkillMatch(BaseModel):
    skills: str


class URLScrape(BaseModel):
    url: str


class SaveJobRequest(BaseModel):
    owner_key: str


class JobMatchRequest(BaseModel):
    skills: str = ""
    experience: str = ""
    education: str = ""


class AdminScrapeSource(BaseModel):
    source: str


class AdminScrapeURL(BaseModel):
    url: str


class CVExportRequest(BaseModel):
    profile: Dict[str, Any]
    template: str = "ats"
    format: str
    include_photo: bool = False
    photo_base64: Optional[str] = None
    target_job: Optional[Dict[str, str]] = None


class CVAIRequest(BaseModel):
    profile: Dict[str, Any]
    job: Optional[Dict[str, Any]] = None
    existing_cv: Optional[Dict[str, Any]] = None
    instructions: str = ""


class UserProfileRequest(BaseModel):
    profile: Dict[str, Any]


class AdvertisementRequest(BaseModel):
    advertisement: Dict[str, Any]


class BusinessRequest(BaseModel):
    profile: Dict[str, Any]


class ReviewRequest(BaseModel):
    action: str
    reason: Optional[str] = None


USER_PROFILE_FIELDS = (
    *CV_PROFILE_FIELDS,
    "photo_data",
    "use_photo",
)


# ============================================================
# ADMIN AUTH
# ============================================================

security_scheme = HTTPBearer(
    auto_error=False
)


def require_admin(
    credentials: Optional[
        HTTPAuthorizationCredentials
    ] = Depends(security_scheme)
):

    configured_token = os.environ.get(
        "ADMIN_API_TOKEN",
        ""
    ).strip()

    if (
        len(configured_token) < 32
        or configured_token.casefold()
        in {
            "replace-with-a-long-random-token",
            "changeme",
            "change-me",
        }
    ):

        raise HTTPException(
            status_code=503,
            detail="Admin access is not configured on this server."
        )

    token = (
        credentials.credentials
        if credentials
        else ""
    )

    if not hmac.compare_digest(
        token,
        configured_token
    ):

        raise HTTPException(
            status_code=401,
            detail="Admin authentication is required.",
            headers={
                "WWW-Authenticate": "Bearer"
            }
        )

    return True


# ============================================================
# INSTANCES
# ============================================================

scraper = JobScraper()
cv_parser = CVParser()
email_sender = EmailSender()


# ============================================================
# SCHEDULER
# ============================================================

scheduler = AsyncIOScheduler()


async def scheduled_auto_scrape():

    if not scrape_lock.acquire(
        blocking=False
    ):
        print(
            "Scheduled scrape skipped because another scrape is running."
        )
        return

    try:

        result = await asyncio.to_thread(
            scrape_all_and_persist
        )

        print(
            f"Scheduled scrape finished: "
            f"{result['added']} new listings, "
            f"{result['duplicates']} duplicates."
        )

    except Exception as e:

        print(
            f"Scheduled scrape failed ({type(e).__name__})."
        )

    finally:

        scrape_lock.release()


# ============================================================
# LIFESPAN
# ============================================================

@asynccontextmanager
async def lifespan(app: FastAPI):

    print("=" * 60)

    print(
        "🚀 Grafan job platform (FastAPI)"
    )

    print("=" * 60)

    print(
        "📖 API documentation: /docs"
    )

    email_enabled = (
        os.environ.get(
            "EMAIL_NOTIFICATIONS",
            "false"
        ).strip().lower()
        in {
            "1",
            "true",
            "yes",
            "on",
        }
        and email_sender.is_configured
    )

    print(
        "📧 Email notifications: "
        f"{'enabled' if email_enabled else 'disabled'}"
    )

    print("=" * 60)

    if (
        os.environ.get(
            "ENABLE_SCRAPE_SCHEDULER",
            "false"
        ).strip().lower()
        in {
            "1",
            "true",
            "yes",
            "on",
        }
    ):

        scheduler.add_job(
            scheduled_auto_scrape,
            "interval",
            hours=6
        )

        scheduler.start()

        print(
            "🕐 Scheduler started (every 6 hours)"
        )

    else:

        print(
            "🕐 Scheduled scraping disabled by configuration"
        )

    print("=" * 60)

    print(
        "✅ Server is running... Press CTRL+C to stop"
    )

    print("=" * 60)

    
    yield

    if scheduler.running:
        scheduler.shutdown()

    print(
        "\n🛑 Server shutting down..."
    )
    
    


# ============================================================
# FASTAPI APP
# ============================================================

app = FastAPI(
    title="Grafan AI Job Platform",
    version="5.0",
    description="AI-Powered Job Platform - Scrape, Match, Apply",
    lifespan=lifespan,
)


app.add_middleware(
    CORSMiddleware,

    allow_origins=[
        origin.strip()
        for origin in os.environ.get(
            "CORS_ORIGINS",
            "http://localhost:3000,"
            "http://127.0.0.1:3000,"
            "http://localhost:3001,"
            "http://127.0.0.1:3001"
        ).split(",")
        if origin.strip()
    ],

    allow_credentials=False,

    allow_methods=[
        "GET",
        "POST",
        "PUT",
        "DELETE",
        "OPTIONS",
    ],

    allow_headers=[
        "Content-Type",
        "Authorization",
        "X-Profile-Key",
    ],
)


# ============================================================
# ROUTES
# ============================================================

@app.get("/")
def home():

    return {
        "name": "Grafan AI Job Platform",
        "version": "5.0",
        "framework": "FastAPI",
        "features": [
            "AI Scraping",
            "CV Builder",
            "Email notifications",
            "Auto-scheduling",
        ],
    }


# ============================================================
# JOB HELPERS
# ============================================================

def is_displayable_job(job):

    if getattr(
        job,
        "job_status",
        "active"
    ) != "active":

        return False

    source = (
        job.source or ""
    ).strip().casefold()

    if source in {
        "sample",
        "demo",
        "mock",
    }:

        return False

    if (
        source in {
            "manual entry",
            "direct post",
            "employer post",
            "public post",
            "custom",
            "grafan",
        }
        or not source
    ):

        return True

    return scraper.is_job(
        job.title or ""
    )


def public_job_payload(job):

    return {
        "id": job.id,
        "title": job.title,
        "company": job.company,
        "location": job.location,
        "job_type": job.job_type,
        "salary": job.salary,
        "description": job.description or None,
        "requirements": job.requirements,
        "job_category": job.job_category,
        "responsibilities": job.responsibilities,
        "qualifications": job.qualifications,
        "skills": [
            item.strip()
            for item in (
                job.required_skills or ""
            ).split(",")
            if item.strip()
        ],
        "experience_requirements": (
            job.experience_requirements
        ),
        "application_deadline": (
            job.application_deadline
        ),
        "application_instructions": (
            job.application_instructions
        ),
        "posted_date": (
            job.posted_date.isoformat()
            if job.posted_date
            else None
        ),
        "application_available": bool(
            job.application_url
            or job.recruiter_email
        ),
        "completeness_score": (
            job.completeness_score or 0
        ),
    }


def admin_job_payload(job):

    return {
        **public_job_payload(job),

        "source": job.source,
        "source_url": job.source_url,
        "application_url": job.application_url,
        "recruiter_email": job.recruiter_email,

        "scraped_at": (
            job.scraped_at.isoformat()
            if job.scraped_at
            else None
        ),

        "job_status": job.job_status,
        "duplicate_of": job.duplicate_of,
    }


# ============================================================
# GET ALL JOBS
# ============================================================

@app.get("/api/jobs")
def get_jobs(
    type: str = "all",
    search: str = "",
    location: str = "",
    limit: int = 24,
    offset: int = 0,
):

    if (
        limit < 1
        or limit > 100
        or offset < 0
    ):

        raise HTTPException(
            status_code=422,
            detail="Invalid job-list page size or offset."
        )

    db = SessionLocal()

    try:

        expire_due_ads(db)

        query = db.query(Job)

        if type != "all":

            query = query.filter_by(
                job_type=type
            )

        if search:

            term = (
                f"%{search.strip()}%"
            )

            query = query.filter(
                Job.title.ilike(term)
                | Job.company.ilike(term)
                | Job.description.ilike(term)
                | Job.location.ilike(term)
            )

        if location:

            query = query.filter(
                Job.location.ilike(
                    f"%{location.strip()}%"
                )
            )

        query = query.filter(
            Job.job_status == "active"
        )

        candidates = (
            query
            .order_by(
                Job.posted_date.desc(),
                Job.id.desc()
            )
            .all()
        )

        visible_jobs = [
            job
            for job in candidates
            if is_displayable_job(job)
        ]

        total = len(
            visible_jobs
        )

        jobs = visible_jobs[
            offset:offset + limit
        ]

        return {
            "success": True,
            "count": total,
            "total": total,
            "limit": limit,
            "offset": offset,
            "jobs": [
                public_job_payload(job)
                for job in jobs
            ],
        }

    finally:

        db.close()


@app.get("/api/jobs/{job_id}")
def get_job(job_id: int):

    db = SessionLocal()

    try:

        job = db.get(
            Job,
            job_id
        )

        if (
            not job
            or not is_displayable_job(job)
        ):

            raise HTTPException(
                status_code=404,
                detail="This job is no longer available."
            )

        return {
            "success": True,
            "job": public_job_payload(job)
        }

    finally:

        db.close()


# ============================================================
# ORIGINAL EMPLOYER APPLICATION LINK
# ============================================================

@app.get("/api/jobs/{job_id}/application")
def get_job_application(
    job_id: int
):

    db = SessionLocal()

    try:

        job = db.get(
            Job,
            job_id
        )

        if (
            not job
            or not is_displayable_job(job)
        ):

            raise HTTPException(
                status_code=404,
                detail="This job is no longer available."
            )

        if job.application_url:

            parsed = urlparse(
                job.application_url
            )

            if (
                parsed.scheme
                in {"http", "https"}
                and parsed.netloc
            ):

                return RedirectResponse(
                    job.application_url,
                    status_code=303
                )

        if (
            job.recruiter_email
            and re.fullmatch(
                r"[^\s@]+@[^\s@]+\.[^\s@]+",
                job.recruiter_email
            )
        ):

            return RedirectResponse(
                f"mailto:{job.recruiter_email}",
                status_code=303
            )

        raise HTTPException(
            status_code=404,
            detail="Application link unavailable"
        )

    finally:

        db.close()


# ============================================================
# SAVE JOB
# ============================================================

@app.post("/api/jobs/{job_id}/save")
def save_job(
    job_id: int,
    data: SaveJobRequest,
):

    owner_key = (
        data.owner_key
        .strip()
        .lower()
    )

    if not re.fullmatch(
        r"[a-f0-9-]{32,36}",
        owner_key
    ):

        raise HTTPException(
            status_code=422,
            detail="Invalid saved-jobs profile key."
        )

    db = SessionLocal()

    try:

        job = db.get(
            Job,
            job_id
        )

        if (
            not job
            or not is_displayable_job(job)
        ):

            raise HTTPException(
                status_code=404,
                detail="This job is no longer available."
            )

        row = (
            db.query(SavedJob)
            .filter_by(
                owner_key=owner_key,
                job_id=job_id
            )
            .first()
        )

        if not row:

            db.add(
                SavedJob(
                    owner_key=owner_key,
                    job_id=job_id
                )
            )

            db.commit()

        return {
            "success": True,
            "saved": True,
            "job_id": job_id,
        }

    finally:

        db.close()


@app.get("/api/saved-jobs")
def get_saved_jobs(
    owner_key: str
):

    owner_key = (
        owner_key
        .strip()
        .lower()
    )

    if not re.fullmatch(
        r"[a-f0-9-]{32,36}",
        owner_key
    ):

        raise HTTPException(
            status_code=422,
            detail="Invalid saved-jobs profile key."
        )

    db = SessionLocal()

    try:

        saved = (
            db.query(SavedJob)
            .filter_by(
                owner_key=owner_key
            )
            .order_by(
                SavedJob.saved_at.desc()
            )
            .all()
        )

        jobs = [
            db.get(Job, row.job_id)
            for row in saved
        ]

        return {
            "success": True,
            "jobs": [
                public_job_payload(job)
                for job in jobs
                if (
                    job
                    and is_displayable_job(job)
                )
            ],
        }

    finally:

        db.close()


# ============================================================
# PROFILE
# ============================================================

def profile_key_hash(
    profile_key
):

    if (
        not profile_key
        or not re.fullmatch(
            r"[a-f0-9]{32}",
            profile_key.strip().lower()
        )
    ):

        raise HTTPException(
            status_code=401,
            detail="A valid browser profile key is required."
        )

    return hashlib.sha256(
        profile_key
        .strip()
        .lower()
        .encode("ascii")
    ).hexdigest()


@app.get("/api/profile")
def get_profile(
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    key_hash = profile_key_hash(
        x_profile_key
    )

    db = SessionLocal()

    try:

        row = db.get(
            UserProfile,
            key_hash
        )

        return {
            "success": True,
            "profile": (
                json.loads(
                    row.profile_json
                )
                if row
                else None
            ),
            "updated_at": (
                row.updated_at.isoformat()
                if row
                else None
            ),
        }

    finally:

        db.close()


@app.post("/api/profile")
def save_profile(
    data: UserProfileRequest,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    key_hash = profile_key_hash(
        x_profile_key
    )

    profile = {
        key: data.profile[key]
        for key in USER_PROFILE_FIELDS
        if key in data.profile
    }

    if any(
        not isinstance(
            value,
            (str, list)
        )
        and not (
            key == "use_photo"
            and isinstance(
                value,
                bool
            )
        )
        for key, value in profile.items()
    ):

        raise HTTPException(
            status_code=422,
            detail="Profile fields must contain text or lists of text."
        )

    if any(
        isinstance(value, list)
        and any(
            not isinstance(
                item,
                str
            )
            for item in value
        )
        for value in profile.values()
    ):

        raise HTTPException(
            status_code=422,
            detail="Profile lists may contain text only."
        )

    if any(
        isinstance(value, str)
        and len(value) > 12000
        for value in profile.values()
    ):

        raise HTTPException(
            status_code=422,
            detail="Each profile field must be 12,000 characters or fewer."
        )

    serialized = json.dumps(
        profile,
        ensure_ascii=False
    )

    if (
        len(
            serialized.encode("utf-8")
        )
        > 5 * 1024 * 1024
    ):

        raise HTTPException(
            status_code=413,
            detail="The profile exceeds the 5 MB limit."
        )

    db = SessionLocal()

    try:

        row = db.get(
            UserProfile,
            key_hash
        )

        if row:

            row.profile_json = serialized
            row.updated_at = datetime.utcnow()

        else:

            db.add(
                UserProfile(
                    key_hash=key_hash,
                    profile_json=serialized,
                    updated_at=datetime.utcnow()
                )
            )

        db.commit()

        return {
            "success": True,
            "saved": True,
            "updated_at": datetime.utcnow().isoformat(),
        }

    finally:

        db.close()


# ============================================================
# UNIVERSAL ADVERTISEMENTS / BUSINESS BRANDING
# ============================================================

AD_CATEGORIES = {
    "jobs",
    "property",
    "land",
    "vehicles",
    "products",
    "services",
    "training-courses",
    "events",
    "business-promotion",
    "ngo-community",
    "freelance",
    "other",
}

AD_TEMPLATES = {
    "modern",
    "professional",
    "corporate",
    "creative",
    "minimal",
    "business",
    "product",
    "property",
    "job",
    "event",
}

AD_CTAS = {
    "Contact Us",
    "WhatsApp Us",
    "Apply Now",
    "Buy Now",
    "Register Now",
    "Learn More",
    "Book Now",
    "Call Now",
    "Visit Website",
}

AD_PROMOTIONS = {
    "free",
    "featured",
    "sponsored",
    "premium",
    "urgent",
}


def ad_owner_hash(
    profile_key
):

    return profile_key_hash(
        profile_key
    )


def safe_ad_payload(
    ad,
    include_private=False
):

    payload = json.loads(
        ad.payload_json or "{}"
    )

    asset_db = SessionLocal()

    try:

        assets = (
            asset_db
            .query(AdvertisementImage)
            .filter_by(
                advertisement_id=ad.id
            )
            .order_by(
                AdvertisementImage.sort_index,
                AdvertisementImage.id
            )
            .all()
        )

        payload["images"] = [
            asset.image_data
            for asset in assets
            if asset.kind == "image"
        ]

        payload["logo"] = next(
            (
                asset.image_data
                for asset in assets
                if asset.kind == "logo"
            ),
            ""
        )

        payload["cover_image"] = next(
            (
                asset.image_data
                for asset in assets
                if asset.kind == "cover"
            ),
            ""
        )

    finally:

        asset_db.close()

    result = {
        "id": ad.id,
        "title": ad.title,
        "category": ad.category,
        "description": ad.description,
        "status": ad.status,
        "rejection_reason": ad.rejection_reason,
        "promotion": ad.promotion,
        "featured": bool(ad.featured),
        "business_id": ad.business_id,
        "job_id": ad.job_id,
        "created_at": (
            ad.created_at.isoformat()
            if ad.created_at
            else None
        ),
        "published_at": (
            ad.published_at.isoformat()
            if ad.published_at
            else None
        ),
        "expires_at": (
            ad.expires_at.isoformat()
            if ad.expires_at
            else None
        ),
        **payload,
    }

    if include_private:

        result["owner_hash"] = ad.owner_hash

    else:

        result.pop(
            "owner_hash",
            None
        )

    return result


def ad_event(
    db,
    ad_id,
    event_type
):

    if event_type not in {
        "view",
        "save",
        "cta",
        "whatsapp",
        "phone",
        "website",
        "application",
        "share",
    }:

        raise HTTPException(
            status_code=422,
            detail="Unsupported analytics event."
        )

    db.add(
        AdvertisementEvent(
            advertisement_id=ad_id,
            event_type=event_type
        )
    )


def save_ad_images(
    db,
    ad_id,
    payload
):

    (
        db.query(AdvertisementImage)
        .filter_by(
            advertisement_id=ad_id
        )
        .delete(
            synchronize_session=False
        )
    )

    for kind, image in (
        ("logo", payload.get("logo")),
        ("cover", payload.get("cover_image")),
    ):

        if image:

            db.add(
                AdvertisementImage(
                    advertisement_id=ad_id,
                    kind=kind,
                    image_data=image
                )
            )

    for index, image in enumerate(
        payload.get("images", [])
    ):

        db.add(
            AdvertisementImage(
                advertisement_id=ad_id,
                kind="image",
                image_data=image,
                sort_index=index
            )
        )


def serialize_ad_content(
    payload
):

    return json.dumps(
        {
            key: value
            for key, value in payload.items()
            if key not in {
                "images",
                "logo",
                "cover_image",
            }
        },
        ensure_ascii=False
    )


def validate_ad_data(
    data
):

    allowed = {
        "title",
        "description",
        "category",
        "location",
        "price",
        "salary",
        "images",
        "contact_phone",
        "whatsapp",
        "email",
        "website",
        "social_links",
        "deadline",
        "details",
        "business_name",
        "logo",
        "cover_image",
        "tagline",
        "brand_primary",
        "brand_secondary",
        "template",
        "cta",
        "cta_url",
        "promotion",
        "expires_at",
        "business_id",
        "job_type",
        "job_sector",
        "required_skills",
        "experience_requirements",
        "qualifications",
        "responsibilities",
        "application_instructions",
    }

    result = {
        key: value
        for key, value in data.items()
        if key in allowed
    }

    title = str(
        result.get(
            "title",
            ""
        )
    ).strip()

    description = str(
        result.get(
            "description",
            ""
        )
    ).strip()

    category = str(
        result.get(
            "category",
            ""
        )
    ).strip().lower()

    if (
        not title
        or len(title) > 200
        or not description
        or len(description) > 12000
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "A title (max 200 characters) "
                "and description (max 12,000 characters) "
                "are required."
            )
        )

    if category not in AD_CATEGORIES:

        raise HTTPException(
            status_code=422,
            detail="Select a valid advertisement category."
        )

    template = result.get(
        "template",
        "modern"
    )

    if template not in AD_TEMPLATES:

        raise HTTPException(
            status_code=422,
            detail="Select a valid branding template."
        )

    cta = result.get(
        "cta",
        "Contact Us"
    )

    if cta not in AD_CTAS:

        raise HTTPException(
            status_code=422,
            detail="Select a supported call-to-action."
        )

    promotion = result.get(
        "promotion",
        "free"
    )

    if promotion not in AD_PROMOTIONS:

        raise HTTPException(
            status_code=422,
            detail="Select a supported promotion type."
        )

    for key in (
        "website",
        "cta_url",
    ):

        value = str(
            result.get(key) or ""
        ).strip()

        if value:

            parsed = urlparse(
                value
            )

            if (
                parsed.scheme
                not in {"http", "https"}
                or not parsed.netloc
            ):

                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"{key.replace('_', ' ').capitalize()} "
                        "must be a valid HTTP or HTTPS URL."
                    )
                )

            result[key] = value

    for key in (
        "email",
    ):

        value = str(
            result.get(key) or ""
        ).strip()

        if value and not re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            value
        ):

            raise HTTPException(
                status_code=422,
                detail="Enter a valid contact email."
            )

    for key in (
        "brand_primary",
        "brand_secondary",
    ):

        value = str(
            result.get(key) or ""
        ).strip()

        if value and not re.fullmatch(
            r"#[0-9a-fA-F]{6}",
            value
        ):

            raise HTTPException(
                status_code=422,
                detail="Brand colors must be six-digit hexadecimal colors."
            )

    social_links = result.get(
        "social_links",
        []
    )

    if (
        not isinstance(
            social_links,
            (list, dict)
        )
        or any(
            not isinstance(value, str)
            for value in (
                social_links.values()
                if isinstance(
                    social_links,
                    dict
                )
                else social_links
            )
        )
    ):

        raise HTTPException(
            status_code=422,
            detail="Social links must be a list or map of text URLs."
        )

    for value in (
        social_links.values()
        if isinstance(
            social_links,
            dict
        )
        else social_links
    ):

        if value and (
            urlparse(value).scheme
            not in {"http", "https"}
            or not urlparse(value).netloc
        ):

            raise HTTPException(
                status_code=422,
                detail="Social links must be valid HTTP or HTTPS URLs."
            )

    expiry_value = str(
        result.get(
            "expires_at"
        ) or ""
    ).strip()

    if expiry_value:

        try:

            datetime.fromisoformat(
                expiry_value.replace(
                    "Z",
                    "+00:00"
                )
            )

        except ValueError:

            raise HTTPException(
                status_code=422,
                detail="Publication expiry must be a valid date or datetime."
            )

    images = result.get(
        "images",
        []
    )

    if (
        not isinstance(
            images,
            list
        )
        or len(images) > 8
    ):

        raise HTTPException(
            status_code=422,
            detail="Add up to eight advertisement images."
        )

    for image in [
        *images,
        result.get("logo"),
        result.get("cover_image"),
    ]:

        if image and (
            not isinstance(image, str)
            or not re.match(
                r"^data:image/(png|jpeg|webp);base64,",
                image
            )
            or len(image) > 3_000_000
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    "Images must be PNG, JPEG, or WebP "
                    "and no larger than 2 MB each."
                )
            )

    result["title"] = title
    result["description"] = description
    result["category"] = category
    result["template"] = template
    result["cta"] = cta

    return result


def ad_is_expired(
    ad
):

    return bool(
        ad.expires_at
        and ad.expires_at
        <= datetime.utcnow()
    )


def parse_ad_expiry(
    value
):

    if not value:
        return None

    parsed = datetime.fromisoformat(
        str(value).replace(
            "Z",
            "+00:00"
        )
    )

    if len(str(value)) == 10:

        return datetime.combine(
            parsed.date(),
            datetime.max.time()
        )

    return (
        parsed.astimezone(
            timezone.utc
        ).replace(
            tzinfo=None
        )
        if parsed.tzinfo
        else parsed
    )


def expire_due_ads(
    db
):

    due = (
        db.query(Advertisement)
        .filter(
            Advertisement.status == "published",
            Advertisement.expires_at.isnot(None),
            Advertisement.expires_at <= datetime.utcnow(),
        )
        .all()
    )

    if not due:
        return

    for ad in due:

        ad.status = "expired"
        ad.updated_at = datetime.utcnow()

        if ad.job_id:

            linked = db.get(
                Job,
                ad.job_id
            )

            if linked:

                linked.job_status = "inactive"

    db.commit()


@app.get("/api/ads/config")
def get_ad_config():

    return {
        "categories": sorted(
            AD_CATEGORIES
        ),
        "templates": sorted(
            AD_TEMPLATES
        ),
        "ctas": sorted(
            AD_CTAS
        ),
        "promotions": sorted(
            AD_PROMOTIONS
        ),
    }


@app.get("/api/business/me")
def get_my_business(
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    db = SessionLocal()

    try:

        row = (
            db.query(BusinessProfile)
            .filter_by(
                owner_hash=owner
            )
            .first()
        )

        return {
            "success": True,
            "business": (
                {
                    "id": row.id,
                    "slug": row.slug,
                    **json.loads(
                        row.profile_json
                    ),
                }
                if row
                else None
            ),
        }

    finally:

        db.close()


@app.post("/api/business/me")
def save_my_business(
    data: BusinessRequest,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    profile = {
        k: v
        for k, v in data.profile.items()
        if k in {
            "name",
            "description",
            "tagline",
            "phone",
            "whatsapp",
            "email",
            "website",
            "linkedin",
            "facebook",
            "instagram",
            "x",
            "location",
            "logo",
            "cover_image",
            "brand_primary",
            "brand_secondary",
        }
    }

    name = str(
        profile.get(
            "name",
            ""
        )
    ).strip()

    if (
        not name
        or len(name) > 160
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Business name is required "
                "(max 160 characters)."
            )
        )

    for key in (
        "brand_primary",
        "brand_secondary",
    ):

        if (
            profile.get(key)
            and not re.fullmatch(
                r"#[0-9a-fA-F]{6}",
                str(profile[key])
            )
        ):

            raise HTTPException(
                status_code=422,
                detail="Brand colors must be six-digit hexadecimal colors."
            )

    for key in (
        "linkedin",
        "facebook",
        "instagram",
        "x",
        "website",
    ):

        value = str(
            profile.get(key) or ""
        ).strip()

        if value and (
            urlparse(value).scheme
            not in {"http", "https"}
            or not urlparse(value).netloc
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    f"{key.capitalize()} "
                    "must be a valid HTTP or HTTPS URL."
                )
            )

    for key in (
        "logo",
        "cover_image",
    ):

        image = profile.get(key)

        if image and (
            not isinstance(
                image,
                str
            )
            or not re.match(
                r"^data:image/(png|jpeg|webp);base64,",
                image
            )
            or len(image) > 3_000_000
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    "Business images must be PNG, JPEG, or WebP "
                    "and no larger than 2 MB each."
                )
            )

    if (
        len(
            json.dumps(
                profile,
                ensure_ascii=False
            ).encode("utf-8")
        )
        > 4_000_000
    ):

        raise HTTPException(
            status_code=413,
            detail="Business profile is too large."
        )

    slug_base = re.sub(
        r"[^a-z0-9]+",
        "-",
        name.casefold()
    ).strip("-")[:80] or "business"

    db = SessionLocal()

    try:

        row = (
            db.query(BusinessProfile)
            .filter_by(
                owner_hash=owner
            )
            .first()
        )

        if row:

            row.profile_json = json.dumps(
                profile,
                ensure_ascii=False
            )

            row.updated_at = datetime.utcnow()

        else:

            slug = slug_base
            suffix = 2

            while (
                db.query(BusinessProfile)
                .filter_by(
                    slug=slug
                )
                .first()
            ):

                slug = (
                    f"{slug_base}-{suffix}"
                )

                suffix += 1

            row = BusinessProfile(
                owner_hash=owner,
                slug=slug,
                profile_json=json.dumps(
                    profile,
                    ensure_ascii=False
                ),
            )

            db.add(row)

        db.commit()
        db.refresh(row)

        return {
            "success": True,
            "business": {
                "id": row.id,
                "slug": row.slug,
                **profile,
            },
        }

    finally:

        db.close()


@app.get("/api/business/{slug}")
def public_business(
    slug: str
):

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        business = (
            db.query(BusinessProfile)
            .filter_by(
                slug=slug
            )
            .first()
        )

        if not business:

            raise HTTPException(
                status_code=404,
                detail="Business page not found."
            )

        ads = (
            db.query(Advertisement)
            .filter_by(
                business_id=business.id,
                status="published"
            )
            .order_by(
                Advertisement.published_at.desc()
            )
            .all()
        )

        ads = [
            ad
            for ad in ads
            if not ad_is_expired(ad)
        ]

        return {
            "success": True,
            "business": {
                "slug": business.slug,
                **json.loads(
                    business.profile_json
                ),
            },
            "advertisements": [
                safe_ad_payload(ad)
                for ad in ads
            ],
        }

    finally:

        db.close()


@app.post("/api/ads", status_code=201)
def create_ad(
    data: AdvertisementRequest,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    ad_data = validate_ad_data(
        data.advertisement
    )

    db = SessionLocal()

    try:

        business_id = ad_data.get(
            "business_id"
        )

        business = (
            db.query(BusinessProfile)
            .filter_by(
                id=business_id,
                owner_hash=owner
            )
            .first()
            if business_id
            else None
        )

        if (
            business_id
            and not business
        ):

            raise HTTPException(
                status_code=404,
                detail="Your business profile was not found."
            )

        if business:

            ad_data[
                "business_slug"
            ] = business.slug

        if ad_data.get(
            "expires_at"
        ):

            ad_expiry = parse_ad_expiry(
                ad_data["expires_at"]
            )

        elif ad_data.get(
            "deadline"
        ):

            try:

                ad_expiry = datetime.combine(
                    datetime.fromisoformat(
                        ad_data["deadline"]
                    ).date(),
                    datetime.max.time()
                )

            except ValueError:

                ad_expiry = None

        else:

            ad_expiry = None

        row = Advertisement(
            owner_hash=owner,
            business_id=(
                business.id
                if business
                else None
            ),
            title=ad_data[
                "title"
            ],
            category=ad_data[
                "category"
            ],
            description=ad_data[
                "description"
            ],
            status="draft",
            payload_json=serialize_ad_content(
                ad_data
            ),
            promotion=ad_data.get(
                "promotion",
                "free"
            ),
            featured=int(
                ad_data.get(
                    "promotion"
                ) == "featured"
            ),
            expires_at=ad_expiry,
        )

        db.add(row)
        db.flush()

        save_ad_images(
            db,
            row.id,
            ad_data
        )

        db.commit()
        db.refresh(row)

        return {
            "success": True,
            "advertisement": safe_ad_payload(
                row
            ),
        }

    finally:

        db.close()


@app.get("/api/ads/mine")
def my_ads(
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        rows = (
            db.query(Advertisement)
            .filter_by(
                owner_hash=owner
            )
            .order_by(
                Advertisement.updated_at.desc()
            )
            .all()
        )

        counts = {
            key: sum(
                ad.status == key
                for ad in rows
            )
            for key in (
                "draft",
                "pending",
                "approved",
                "published",
                "rejected",
                "expired",
                "closed",
                "suspended",
            )
        }

        for ad in rows:

            if (
                ad.status == "published"
                and ad_is_expired(ad)
            ):

                counts["expired"] += 1

        output = []

        for ad in rows:

            item = safe_ad_payload(
                ad
            )

            item["analytics"] = {
                e: db.query(
                    AdvertisementEvent
                )
                .filter_by(
                    advertisement_id=ad.id,
                    event_type=e
                )
                .count()
                for e in (
                    "view",
                    "save",
                    "cta",
                    "whatsapp",
                    "phone",
                    "website",
                    "application",
                    "share",
                )
            }

            output.append(
                item
            )

        return {
            "success": True,
            "counts": counts,
            "advertisements": output,
        }

    finally:

        db.close()


@app.put("/api/ads/{ad_id}")
def update_my_ad(
    ad_id: int,
    data: AdvertisementRequest,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    updated = validate_ad_data(
        data.advertisement
    )

    db = SessionLocal()

    try:

        row = (
            db.query(Advertisement)
            .filter_by(
                id=ad_id,
                owner_hash=owner
            )
            .first()
        )

        if not row:

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        if row.status not in {
            "draft",
            "rejected",
        }:

            raise HTTPException(
                status_code=409,
                detail=(
                    "Only drafts and rejected "
                    "advertisements can be edited."
                )
            )

        business_id = updated.get(
            "business_id"
        )

        business = (
            db.query(BusinessProfile)
            .filter_by(
                id=business_id,
                owner_hash=owner
            )
            .first()
            if business_id
            else None
        )

        if (
            business_id
            and not business
        ):

            raise HTTPException(
                status_code=404,
                detail="Your business profile was not found."
            )

        if business:

            updated[
                "business_slug"
            ] = business.slug

        if updated.get(
            "expires_at"
        ):

            row.expires_at = parse_ad_expiry(
                updated["expires_at"]
            )

        elif updated.get(
            "deadline"
        ):

            try:

                row.expires_at = datetime.combine(
                    datetime.fromisoformat(
                        updated["deadline"]
                    ).date(),
                    datetime.max.time()
                )

            except ValueError:

                row.expires_at = None

        else:

            row.expires_at = None

        row.title = updated[
            "title"
        ]

        row.category = updated[
            "category"
        ]

        row.description = updated[
            "description"
        ]

        row.business_id = (
            business.id
            if business
            else None
        )

        row.payload_json = serialize_ad_content(
            updated
        )

        row.status = "draft"

        row.rejection_reason = None

        row.updated_at = datetime.utcnow()

        save_ad_images(
            db,
            row.id,
            updated
        )

        db.commit()

        return {
            "success": True,
            "advertisement": safe_ad_payload(
                row
            ),
        }

    finally:

        db.close()


@app.post("/api/ads/{ad_id}/submit")
def submit_my_ad(
    ad_id: int,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    db = SessionLocal()

    try:

        row = (
            db.query(Advertisement)
            .filter_by(
                id=ad_id,
                owner_hash=owner
            )
            .first()
        )

        if not row:

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        if row.status not in {
            "draft",
            "rejected",
        }:

            raise HTTPException(
                status_code=409,
                detail=(
                    "This advertisement cannot "
                    "be submitted in its current state."
                )
            )

        row.status = "pending"
        row.rejection_reason = None
        row.updated_at = datetime.utcnow()

        db.commit()

        return {
            "success": True,
            "advertisement": safe_ad_payload(
                row
            ),
        }

    finally:

        db.close()


@app.post("/api/ads/{ad_id}/actions/{action}")
def act_on_my_ad(
    ad_id: int,
    action: str,
    x_profile_key: Optional[str] = Header(
        default=None
    )
):

    owner = ad_owner_hash(
        x_profile_key
    )

    db = SessionLocal()

    try:

        row = (
            db.query(Advertisement)
            .filter_by(
                id=ad_id,
                owner_hash=owner
            )
            .first()
        )

        if not row:

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        if (
            action == "close"
            and row.status == "published"
        ):

            row.status = "closed"

            if row.job_id:

                linked = db.get(
                    Job,
                    row.job_id
                )

                if linked:

                    linked.job_status = "inactive"

        elif action == "duplicate":

            clone = Advertisement(
                owner_hash=owner,
                business_id=row.business_id,
                title=row.title,
                category=row.category,
                description=row.description,
                status="draft",
                payload_json=row.payload_json,
                promotion="free",
                featured=0,
            )

            db.add(
                clone
            )

            db.flush()

            for asset in (
                db.query(
                    AdvertisementImage
                )
                .filter_by(
                    advertisement_id=row.id
                )
                .all()
            ):

                db.add(
                    AdvertisementImage(
                        advertisement_id=clone.id,
                        kind=asset.kind,
                        image_data=asset.image_data,
                        sort_index=asset.sort_index,
                    )
                )

        elif action == "delete":

            if row.job_id:

                linked = db.get(
                    Job,
                    row.job_id
                )

                if linked:

                    db.delete(
                        linked
                    )

            (
                db.query(
                    AdvertisementEvent
                )
                .filter_by(
                    advertisement_id=row.id
                )
                .delete(
                    synchronize_session=False
                )
            )

            (
                db.query(
                    AdvertisementImage
                )
                .filter_by(
                    advertisement_id=row.id
                )
                .delete(
                    synchronize_session=False
                )
            )

            db.delete(
                row
            )

        else:

            raise HTTPException(
                status_code=409,
                detail=(
                    "That action is unavailable "
                    "for this advertisement."
                )
            )

        db.commit()

        return {
            "success": True
        }

    finally:

        db.close()


@app.get("/api/ads")
def public_ads(
    category: str = "",
    search: str = "",
    limit: int = 30,
    offset: int = 0,
):

    if (
        limit < 1
        or limit > 100
        or offset < 0
    ):

        raise HTTPException(
            status_code=422,
            detail="Invalid page size or offset."
        )

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        q = (
            db.query(
                Advertisement
            )
            .filter_by(
                status="published"
            )
        )

        if category:

            q = q.filter_by(
                category=category
            )

        if search.strip():

            q = q.filter(
                Advertisement.title.ilike(
                    f"%{search.strip()}%"
                )
                | Advertisement.description.ilike(
                    f"%{search.strip()}%"
                )
            )

        rows = [
            ad
            for ad in (
                q.order_by(
                    Advertisement.featured.desc(),
                    Advertisement.published_at.desc(),
                )
                .all()
            )
            if not ad_is_expired(ad)
        ]

        return {
            "success": True,
            "total": len(rows),
            "advertisements": [
                safe_ad_payload(ad)
                for ad in rows[
                    offset:offset + limit
                ]
            ],
        }

    finally:

        db.close()


@app.get("/api/ads/{ad_id}")
def public_ad(
    ad_id: int
):

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        row = db.get(
            Advertisement,
            ad_id
        )

        if (
            not row
            or row.status != "published"
            or ad_is_expired(row)
        ):

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        ad_event(
            db,
            row.id,
            "view"
        )

        db.commit()

        return {
            "success": True,
            "advertisement": safe_ad_payload(
                row
            ),
        }

    finally:

        db.close()


@app.post("/api/ads/{ad_id}/track/{event_type}")
def track_ad_event(
    ad_id: int,
    event_type: str
):

    db = SessionLocal()

    try:

        row = db.get(
            Advertisement,
            ad_id
        )

        if (
            not row
            or row.status != "published"
            or ad_is_expired(row)
        ):

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        ad_event(
            db,
            row.id,
            event_type
        )

        db.commit()

        return {
            "success": True
        }

    finally:

        db.close()


@app.get("/api/admin/ads")
def admin_ads(
    category: str = "",
    status: str = "",
    search: str = "",
    _admin: bool = Depends(require_admin),
):

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        q = db.query(
            Advertisement
        )

        if category:

            q = q.filter_by(
                category=category
            )

        if status:

            q = q.filter_by(
                status=status
            )

        if search:

            q = q.filter(
                Advertisement.title.ilike(
                    f"%{search}%"
                )
                | Advertisement.description.ilike(
                    f"%{search}%"
                )
            )

        rows = (
            q.order_by(
                Advertisement.created_at.desc()
            )
            .all()
        )

        counts = {
            s: db.query(
                Advertisement
            )
            .filter_by(
                status=s
            )
            .count()
            for s in (
                "pending",
                "published",
                "rejected",
                "expired",
            )
        }

        counts["total"] = (
            db.query(
                Advertisement
            ).count()
        )

        counts["featured"] = (
            db.query(
                Advertisement
            )
            .filter_by(
                featured=1
            )
            .count()
        )

        counts["advertisers"] = (
            db.query(
                Advertisement.owner_hash
            )
            .distinct()
            .count()
        )

        counts["businesses"] = (
            db.query(
                BusinessProfile
            )
            .count()
        )

        return {
            "success": True,
            "counts": counts,
            "advertisements": [
                safe_ad_payload(
                    r,
                    True
                )
                for r in rows
            ],
        }

    finally:

        db.close()


@app.post("/api/admin/ads/{ad_id}/review")
def admin_review_ad(
    ad_id: int,
    data: ReviewRequest,
    _admin: bool = Depends(require_admin),
):

    db = SessionLocal()

    try:

        row = db.get(
            Advertisement,
            ad_id
        )

        if not row:

            raise HTTPException(
                status_code=404,
                detail="Advertisement not found."
            )

        action = data.action

        if (
            action == "approve"
            and row.status == "pending"
        ):

            row.status = "approved"

        elif (
            action == "publish"
            and row.status == "approved"
        ):

            row.status = "published"
            row.published_at = datetime.utcnow()

        elif (
            action == "reject"
            and row.status == "pending"
        ):

            if not (
                data.reason
                or ""
            ).strip():

                raise HTTPException(
                    status_code=422,
                    detail="A rejection reason is required."
                )

            row.status = "rejected"

            row.rejection_reason = (
                data.reason.strip()[:2000]
            )

        elif (
            action == "request_changes"
            and row.status == "pending"
        ):

            if not (
                data.reason
                or ""
            ).strip():

                raise HTTPException(
                    status_code=422,
                    detail="A change request reason is required."
                )

            row.status = "rejected"

            row.rejection_reason = (
                "Changes requested: "
                + data.reason.strip()[:1900]
            )

        elif (
            action == "suspend"
            and row.status == "published"
        ):

            row.status = "suspended"

            if row.job_id:

                linked = db.get(
                    Job,
                    row.job_id
                )

                if linked:

                    linked.job_status = "inactive"

        elif (
            action == "restore"
            and row.status == "suspended"
        ):

            row.status = "approved"

            if row.job_id:

                linked = db.get(
                    Job,
                    row.job_id
                )

                if linked:

                    linked.job_status = "inactive"

        elif action == "delete":

            if row.job_id:

                linked = db.get(
                    Job,
                    row.job_id
                )

                if linked:

                    db.delete(
                        linked
                    )

            (
                db.query(
                    AdvertisementEvent
                )
                .filter_by(
                    advertisement_id=row.id
                )
                .delete(
                    synchronize_session=False
                )
            )

            (
                db.query(
                    AdvertisementImage
                )
                .filter_by(
                    advertisement_id=row.id
                )
                .delete(
                    synchronize_session=False
                )
            )

            db.delete(
                row
            )

            db.commit()

            return {
                "success": True
            }

        else:

            raise HTTPException(
                status_code=409,
                detail=(
                    "That review action is not valid "
                    "in the current state."
                )
            )

        if (
            action == "publish"
            and row.category == "jobs"
            and not row.job_id
        ):

            payload = json.loads(
                row.payload_json
                or "{}"
            )

            job_type = (
                payload.get(
                    "job_type"
                )
                or "onsite"
            )

            if job_type not in {
                "remote",
                "hybrid",
                "onsite",
                "gig",
            }:

                job_type = "onsite"

            job = Job(
                title=row.title,

                company=(
                    payload.get(
                        "business_name"
                    )
                    or "Employer"
                )[:100],

                job_type=job_type,

                job_category=payload.get(
                    "job_sector"
                ),

                location=payload.get(
                    "location"
                ),

                salary=(
                    payload.get(
                        "salary"
                    )
                    or payload.get(
                        "price"
                    )
                ),

                description=row.description,

                source="Employer Post",

                application_url=(
                    payload.get(
                        "cta_url"
                    )
                    or (
                        payload.get(
                            "website"
                        )
                        if payload.get(
                            "cta"
                        ) == "Visit Website"
                        else ""
                    )
                ),

                recruiter_email=payload.get(
                    "email"
                ),

                application_deadline=payload.get(
                    "deadline"
                ),

                application_instructions=payload.get(
                    "application_instructions"
                ),

                requirements=payload.get(
                    "details"
                ),

                responsibilities=payload.get(
                    "responsibilities"
                ),

                qualifications=payload.get(
                    "qualifications"
                ),

                required_skills=payload.get(
                    "required_skills"
                ),

                experience_requirements=payload.get(
                    "experience_requirements"
                ),

                job_status="active",

                posted_date=datetime.utcnow()
            )

            db.add(
                job
            )

            db.flush()

            row.job_id = job.id

        row.updated_at = datetime.utcnow()

        db.commit()

        return {
            "success": True,
            "advertisement": safe_ad_payload(
                row,
                True
            ),
        }

    finally:

        db.close()


# ============================================================
# JOB MATCHING - SINGLE JOB
# ============================================================

@app.post("/api/jobs/{job_id}/match")
def match_one_job(
    job_id: int,
    data: JobMatchRequest
):

    skills = [
        part.strip()
        for part in data.skills.split(",")
        if part.strip()
    ]

    if (
        len(data.skills) > 1000
        or len(data.experience) > 2000
        or len(data.education) > 2000
    ):

        raise HTTPException(
            status_code=422,
            detail="Match details are too long."
        )

    db = SessionLocal()

    try:

        job = db.get(
            Job,
            job_id
        )

        if (
            not job
            or not is_displayable_job(job)
        ):

            raise HTTPException(
                status_code=404,
                detail="This job is no longer available."
            )

        job_text = " ".join(
            filter(
                None,
                [
                    job.title,
                    job.description,
                    job.requirements,
                    job.responsibilities,
                    job.qualifications,
                    job.required_skills,
                ]
            )
        ).casefold()

        matched = [
            skill
            for skill in skills
            if re.search(
                r"(?<![a-z0-9])"
                + re.escape(
                    skill.casefold()
                )
                + r"(?![a-z0-9])",
                job_text
            )
        ]

        missing = [
            skill
            for skill in skills
            if skill not in matched
        ]

        skills_score = (
            round(
                100
                * len(matched)
                / len(skills)
            )
            if skills
            else None
        )

        exp_score = None
        edu_score = None

        available_scores = [
            s
            for s in (
                skills_score,
                exp_score,
                edu_score
            )
            if s is not None
        ]

        overall = (
            round(
                sum(available_scores)
                / len(available_scores)
            )
            if available_scores
            else None
        )

        return {
            "success": True,
            "job_id": job_id,
            "skills_match": skills_score,
            "experience_match": exp_score,
            "education_match": edu_score,
            "overall_match": overall,
            "matched_skills": matched,
            "missing_skills": missing,
            "explanation": (
                "Scores compare only information supplied by "
                "you with fields explicitly present in this listing. "
                "Missing listing fields are not scored."
            ),
        }

    finally:

        db.close()


# ============================================================
# CREATE JOB
# ============================================================

@app.post("/api/jobs", status_code=201)
def create_job(
    job: JobCreate,
    _admin: bool = Depends(require_admin)
):

    title = job.title.strip()

    company = (
        job.company or ""
    ).strip()

    if not title:

        raise HTTPException(
            status_code=422,
            detail="Job title is required."
        )

    if not company:

        raise HTTPException(
            status_code=422,
            detail="Company name is required."
        )

    if (
        job.job_type is not None
        and job.job_type not in {
            "remote",
            "hybrid",
            "onsite",
            "gig",
        }
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Job type must be "
                "remote, hybrid, onsite, or gig."
            )
        )

    if (
        len(title) > 200
        or len(company) > 100
    ):

        raise HTTPException(
            status_code=422,
            detail="Job title or company name is too long."
        )

    for field_name, value in (
        ("Location", job.location),
        ("Salary", job.salary),
    ):

        if (
            value
            and len(value.strip()) > 150
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    f"{field_name} is too long."
                )
            )

    field_limits = {
        "description": (
            job.description,
            12000
        ),
        "requirements": (
            job.requirements,
            8000
        ),
        "responsibilities": (
            job.responsibilities,
            8000
        ),
        "qualifications": (
            job.qualifications,
            8000
        ),
        "required_skills": (
            job.required_skills,
            2000
        ),
        "experience_requirements": (
            job.experience_requirements,
            2000
        ),
        "application_instructions": (
            job.application_instructions,
            4000
        ),
        "application_deadline": (
            job.application_deadline,
            100
        ),
    }

    for field, (
        value,
        maximum
    ) in field_limits.items():

        if value and len(value) > maximum:

            raise HTTPException(
                status_code=422,
                detail=(
                    f"{field.replace('_', ' ').capitalize()} "
                    f"must be {maximum} characters or fewer."
                )
            )

    if (
        not (
            job.application_url
            or ""
        ).strip()
        and not (
            job.recruiter_email
            or ""
        ).strip()
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Provide an application URL "
                "or recruiter email."
            )
        )

    for label, value in (
        (
            "Application URL",
            job.application_url
        ),
        (
            "Source URL",
            job.source_url
        ),
    ):

        if value:

            parsed = urlparse(
                value.strip()
            )

            if (
                parsed.scheme
                not in {"http", "https"}
                or not parsed.netloc
            ):

                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"{label} must be a valid "
                        "HTTP or HTTPS link."
                    )
                )

    if (
        job.recruiter_email
        and not re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            job.recruiter_email.strip()
        )
    ):

        raise HTTPException(
            status_code=422,
            detail="Enter a valid recruiter email address."
        )

    db = SessionLocal()

    try:

        new_job = Job(
            title=title,
            company=company,
            job_type=job.job_type,
            location=job.location,
            salary=job.salary,
            description=job.description,
            source=job.source,
            source_url=(
                job.source_url or ""
            ).strip()
            or None,
            application_url=(
                job.application_url
                or ""
            ).strip()
            or None,
            recruiter_email=job.recruiter_email,
            requirements=job.requirements,
            responsibilities=job.responsibilities,
            qualifications=job.qualifications,
            required_skills=job.required_skills,
            experience_requirements=job.experience_requirements,
            application_deadline=job.application_deadline,
            application_instructions=job.application_instructions,

            completeness_score=sum(
                bool(value)
                for value in [
                    title,
                    company,
                    job.location,
                    job.job_type,
                    job.salary,
                    job.description,
                    job.requirements,
                    job.responsibilities,
                    job.qualifications,
                    job.required_skills,
                    job.application_url,
                    job.application_instructions,
                ]
            ),

            posted_date=datetime.now()
        )

        db.add(
            new_job
        )

        db.commit()

        db.refresh(
            new_job
        )

        print(
            f"✅ New job: "
            f"{new_job.title} "
            f"at "
            f"{new_job.company}"
        )

        return {
            "success": True,
            "message": "Job posted successfully!",
            "job": {
                "id": new_job.id,
                "title": new_job.title,
                "company": new_job.company,
                "job_type": new_job.job_type,
                "location": new_job.location,
                "salary": new_job.salary,
            },
        }

    finally:

        db.close()


# ============================================================
# DELETE JOB
# ============================================================

@app.delete("/api/jobs/{job_id}")
def delete_job(
    job_id: int,
    _admin: bool = Depends(require_admin)
):

    db = SessionLocal()

    try:

        job = db.query(
            Job
        ).get(
            job_id
        )

        if not job:

            raise HTTPException(
                status_code=404,
                detail="Job not found"
            )

        title = job.title

        db.delete(
            job
        )

        db.commit()

        return {
            "success": True,
            "message": (
                f'Job "{title}" '
                "deleted successfully"
            ),
        }

    finally:

        db.close()


# ============================================================
# UPDATE JOB
# ============================================================

@app.put("/api/jobs/{job_id}")
def update_job(
    job_id: int,
    data: JobUpdate,
    _admin: bool = Depends(require_admin)
):

    changes = data.dict(
        exclude_unset=True
    )

    if (
        "title" in changes
        and changes["title"] is not None
        and not changes["title"].strip()
    ):

        raise HTTPException(
            status_code=422,
            detail="Job title is required."
        )

    if (
        "job_status" in changes
        and changes["job_status"]
        not in {
            "active",
            "inactive",
        }
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Job status must be "
                "active or inactive."
            )
        )

    if (
        changes.get("job_type")
        is not None
        and changes.get("job_type")
        not in {
            "remote",
            "hybrid",
            "onsite",
            "gig",
        }
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Job type must be "
                "remote, hybrid, onsite, or gig."
            )
        )

    for field, maximum in (
        (
            "title",
            200
        ),
        (
            "company",
            100
        ),
        (
            "location",
            150
        ),
        (
            "salary",
            150
        ),
        (
            "description",
            12000
        ),
        (
            "requirements",
            8000
        ),
        (
            "responsibilities",
            8000
        ),
        (
            "qualifications",
            8000
        ),
        (
            "required_skills",
            2000
        ),
        (
            "experience_requirements",
            2000
        ),
        (
            "application_instructions",
            4000
        ),
        (
            "application_deadline",
            100
        ),
        (
            "recruiter_email",
            100
        ),
    ):

        if (
            changes.get(field)
            is not None
            and len(
                changes[field].strip()
            ) > maximum
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    f"{field.replace('_', ' ').capitalize()} "
                    f"is too long."
                )
            )

    for field in (
        "application_url",
        "source_url",
    ):

        value = changes.get(
            field
        )

        if value:

            parsed = urlparse(
                value.strip()
            )

            if (
                parsed.scheme
                not in {
                    "http",
                    "https",
                }
                or not parsed.netloc
            ):

                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"{field.replace('_', ' ').capitalize()} "
                        "must be a valid HTTP or HTTPS link."
                    )
                )

    if (
        changes.get(
            "recruiter_email"
        )
        and not re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            changes[
                "recruiter_email"
            ].strip()
        )
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Enter a valid recruiter email address."
            )
        )

    db = SessionLocal()

    try:

        job = db.get(
            Job,
            job_id
        )

        if not job:

            raise HTTPException(
                status_code=404,
                detail="Job not found"
            )

        for field, value in changes.items():

            if value is not None:

                setattr(
                    job,
                    field,
                    value
                )

        job.completeness_score = sum(
            bool(value)
            for value in [
                job.title,
                job.company,
                job.location,
                job.job_type,
                job.salary,
                job.description,
                job.requirements,
                job.responsibilities,
                job.qualifications,
                job.required_skills,
                job.application_url,
                job.application_instructions,
            ]
        )

        db.commit()
        db.refresh(job)

        return {
            "success": True,
            "message": "Job updated successfully!",
            "job": {
                "id": job.id,
                "title": job.title,
                "company": job.company,
                "job_type": job.job_type,
                "location": job.location,
                "salary": job.salary,
            },
        }

    finally:

        db.close()


# ============================================================
# SCRAPING
# ============================================================

def parse_scraped_datetime(
    value
):

    if not value:
        return None

    try:

        parsed = (
            datetime.fromisoformat(
                value.replace(
                    "Z",
                    "+00:00"
                )
            )
            if isinstance(
                value,
                str
            )
            else value
        )

        return (
            parsed.replace(
                tzinfo=None
            )
            if parsed.tzinfo
            else parsed
        )

    except (
        TypeError,
        ValueError
    ):

        return None


def find_duplicate_job(
    db,
    item
):

    title = (
        item.get(
            "title"
        )
        or ""
    ).strip().casefold()

    company = (
        item.get(
            "company"
        )
        or ""
    ).strip().casefold()

    source = (
        item.get(
            "source"
        )
        or "Scraped"
    ).strip().casefold()

    candidates = (
        db.query(Job)
        .filter(
            func.lower(
                func.trim(
                    Job.title
                )
            ) == title,

            func.lower(
                func.trim(
                    func.coalesce(
                        Job.company,
                        ""
                    )
                )
            ) == company,
        )
    )

    source_url = (
        item.get(
            "source_url"
        )
        or ""
    ).strip()

    if source_url:

        same_url = (
            candidates
            .filter(
                Job.source_url
                == source_url
            )
            .first()
        )

        if same_url:

            return same_url

    if not company:

        return (
            candidates
            .filter(
                func.lower(
                    func.coalesce(
                        Job.source,
                        ""
                    )
                )
                == source
            )
            .first()
        )

    return (
        candidates
        .order_by(
            Job.id.asc()
        )
        .first()
    )


def validated_scrape_application_fields(
    item
):

    application_url = (
        item.get(
            "application_url"
        )
        or ""
    ).strip()

    parsed = (
        urlparse(
            application_url
        )
        if application_url
        else None
    )

    if (
        not parsed
        or parsed.scheme
        not in {
            "http",
            "https",
        }
        or not parsed.netloc
    ):

        application_url = None

    recruiter_email = (
        item.get(
            "recruiter_email"
        )
        or ""
    ).strip()

    if not re.fullmatch(
        r"[^\s@]+@[^\s@]+\.[^\s@]+",
        recruiter_email
    ):

        recruiter_email = None

    return (
        application_url,
        recruiter_email
    )


def scrape_all_and_persist(
    progress_callback=None
):

    db = SessionLocal()

    try:

        if progress_callback:

            progress_callback(
                "Fetching listings from configured sources"
            )

        jobs = scraper.scrape_all_sources()

        if progress_callback:

            progress_callback(
                "Validating listings and checking application links"
            )

        new_count = 0
        duplicate_count = 0
        invalid_count = 0
        linked_count = 0

        for jd in jobs:

            title = (
                jd.get(
                    "title"
                )
                if isinstance(
                    jd,
                    dict
                )
                else None
            )

            if (
                not isinstance(
                    title,
                    str
                )
                or not title.strip()
                or len(
                    title.strip()
                ) > 200
                or not scraper.is_job(
                    title
                )
            ):

                invalid_count += 1

                continue

            existing = find_duplicate_job(
                db,
                jd
            )

            if not existing:

                (
                    application_url,
                    recruiter_email
                ) = (
                    validated_scrape_application_fields(
                        jd
                    )
                )

                row = Job(

                    title=title.strip(),

                    company=jd.get(
                        "company"
                    ),

                    job_type=jd.get(
                        "job_type"
                    ),

                    location=jd.get(
                        "location"
                    ),

                    salary=jd.get(
                        "salary"
                    ),

                    description=jd.get(
                        "description"
                    ),

                    source=jd.get(
                        "source",
                        "AI-Scraped"
                    ),

                    source_url=(
                        jd.get(
                            "source_url"
                        )
                        or jd.get(
                            "url"
                        )
                    ),

                    scraped_at=(
                        parse_scraped_datetime(
                            jd.get(
                                "scraped_at"
                            )
                        )
                        or datetime.utcnow()
                    ),

                    application_url=application_url,

                    recruiter_email=recruiter_email,

                    requirements=jd.get(
                        "requirements"
                    ),

                    responsibilities=jd.get(
                        "responsibilities"
                    ),

                    qualifications=jd.get(
                        "qualifications"
                    ),

                    required_skills=jd.get(
                        "required_skills"
                    ),

                    experience_requirements=jd.get(
                        "experience_requirements"
                    ),

                    application_deadline=jd.get(
                        "application_deadline"
                    ),

                    application_instructions=jd.get(
                        "application_instructions"
                    ),

                    posted_date=(
                        parse_scraped_datetime(
                            jd.get(
                                "posted_date"
                            )
                        )
                    ),
                )

                row.completeness_score = sum(
                    bool(value)
                    for value in [
                        row.title,
                        row.company,
                        row.location,
                        row.job_type,
                        row.salary,
                        row.description,
                        row.requirements,
                        row.responsibilities,
                        row.qualifications,
                        row.required_skills,
                        row.application_url,
                    ]
                )

                db.add(
                    row
                )

                new_count += 1

                linked_count += bool(
                    row.application_url
                    or row.recruiter_email
                )

            else:

                duplicate_count += 1

        if progress_callback:

            progress_callback(
                "Saving validated listings"
            )

        db.commit()

        return {
            "discovered": len(jobs),
            "added": new_count,
            "duplicates": duplicate_count,
            "invalid": invalid_count,
            "with_application_links": linked_count,
            "without_application_links": max(
                0,
                new_count - linked_count
            ),
            "errors": scraper.last_errors,
        }

    finally:

        db.close()


def set_scrape_progress(
    task_id: str,
    progress: str
):

    scrape_runs[
        task_id
    ] = {
        **scrape_runs.get(
            task_id,
            {}
        ),
        "state": "running",
        "progress": progress,
    }


async def run_scrape_task(
    task_id: str,
    url: Optional[str] = None
):

    try:

        progress = (
            lambda step:
            set_scrape_progress(
                task_id,
                step
            )
        )

        if url:

            result = await asyncio.to_thread(
                scrape_url_and_persist,
                url,
                progress
            )

        else:

            result = await asyncio.to_thread(
                scrape_all_and_persist,
                progress
            )

        if (
            result.get("errors")
            and not result.get("discovered")
        ):

            scrape_runs[
                task_id
            ] = {
                "state": "failed",
                "error": "No sources could be scraped.",
                "source_errors": result[
                    "errors"
                ],
            }

        else:

            scrape_runs[
                task_id
            ] = {
                "state": "completed",
                "result": result,
            }

    except Exception as e:

        print(
            f"Scrape task failed ({type(e).__name__})."
        )

        scrape_runs[
            task_id
        ] = {
            "state": "failed",
            "error": (
                "The website could not be reached "
                "or its listing page could not be read."
            ),
        }

    finally:

        scrape_lock.release()


async def run_source_scrape_task(
    task_id: str,
    source_name: str,
    method_name: str
):

    try:

        set_scrape_progress(
            task_id,
            "Fetching listings"
        )

        jobs = await asyncio.to_thread(
            getattr(
                scraper,
                method_name
            )
        )

        set_scrape_progress(
            task_id,
            "Validating listings, links, and duplicates"
        )

        db = SessionLocal()

        added = 0
        duplicates = 0
        linked = 0

        try:

            for item in jobs:

                if (
                    not item.get(
                        "title"
                    )
                    or not scraper.is_job(
                        item.get(
                            "title"
                        )
                    )
                ):

                    continue

                if find_duplicate_job(
                    db,
                    item
                ):

                    duplicates += 1

                    continue

                (
                    application_url,
                    recruiter_email
                ) = (
                    validated_scrape_application_fields(
                        item
                    )
                )

                row = Job(
                    title=item[
                        "title"
                    ],

                    company=item.get(
                        "company"
                    ),

                    job_type=item.get(
                        "job_type"
                    ),

                    location=item.get(
                        "location"
                    ),

                    salary=item.get(
                        "salary"
                    ),

                    description=item.get(
                        "description"
                    ),

                    source=source_name,

                    source_url=(
                        item.get(
                            "source_url"
                        )
                        or item.get(
                            "url"
                        )
                    ),

                    application_url=application_url,

                    requirements=item.get(
                        "requirements"
                    ),

                    responsibilities=item.get(
                        "responsibilities"
                    ),

                    qualifications=item.get(
                        "qualifications"
                    ),

                    required_skills=item.get(
                        "required_skills"
                    ),

                    experience_requirements=item.get(
                        "experience_requirements"
                    ),

                    application_deadline=item.get(
                        "application_deadline"
                    ),

                    application_instructions=item.get(
                        "application_instructions"
                    ),

                    recruiter_email=recruiter_email,

                    scraped_at=(
                        parse_scraped_datetime(
                            item.get(
                                "scraped_at"
                            )
                        )
                        or datetime.utcnow()
                    ),

                    posted_date=(
                        parse_scraped_datetime(
                            item.get(
                                "posted_date"
                            )
                        )
                    ),
                )

                row.completeness_score = sum(
                    bool(value)
                    for value in [
                        row.title,
                        row.company,
                        row.location,
                        row.job_type,
                        row.salary,
                        row.description,
                        row.requirements,
                        row.responsibilities,
                        row.qualifications,
                        row.required_skills,
                        row.application_url,
                    ]
                )

                db.add(
                    row
                )

                added += 1

                linked += bool(
                    row.application_url
                    or row.recruiter_email
                )

            set_scrape_progress(
                task_id,
                "Saving validated listings"
            )

            db.commit()

            scrape_runs[
                task_id
            ] = {
                "state": "completed",
                "result": {
                    "source": source_name,
                    "discovered": len(jobs),
                    "added": added,
                    "duplicates": duplicates,
                    "invalid": max(
                        0,
                        len(jobs)
                        - added
                        - duplicates
                    ),
                    "with_application_links": linked,
                    "without_application_links": max(
                        0,
                        added - linked
                    ),
                },
            }

        finally:

            db.close()

    except Exception as error:

        print(
            f"Source scrape failed for "
            f"{source_name} "
            f"({type(error).__name__})."
        )

        scrape_runs[
            task_id
        ] = {
            "state": "failed",
            "error": (
                "The source could not be reached "
                "or its listings could not be read."
            ),
        }

    finally:

        scrape_lock.release()


@app.post("/api/scrape")
def trigger_scrape(
    background_tasks: BackgroundTasks,
    _admin: bool = Depends(require_admin),
    url: Optional[str] = None
):

    if not scrape_lock.acquire(
        blocking=False
    ):

        raise HTTPException(
            status_code=409,
            detail=(
                "A scraping request is "
                "already running."
            )
        )

    task_id = uuid.uuid4().hex

    scrape_runs[
        task_id
    ] = {
        "state": "running",
        "progress": "Fetching listings",
    }

    background_tasks.add_task(
        run_scrape_task,
        task_id,
        url
    )

    return {
        "success": True,
        "task_id": task_id,
        "state": "running",
    }


@app.get("/api/scrape/status/{task_id}")
def get_scrape_status(
    task_id: str,
    _admin: bool = Depends(require_admin)
):

    status = scrape_runs.get(
        task_id
    )

    if not status:

        raise HTTPException(
            status_code=404,
            detail=(
                "Scrape task was not found "
                "or has expired."
            )
        )

    return status


# ============================================================
# CUSTOM URL SCRAPE
# ============================================================

def scrape_url_and_persist(
    url: str,
    progress_callback=None
):

    db = SessionLocal()

    try:

        jobs = scraper.scrape_custom_url(
            url
        )

        if progress_callback:

            progress_callback(
                "Validating listings, links, and duplicates"
            )

        new_count = 0
        duplicate_count = 0
        invalid_count = 0
        linked_count = 0

        for jd in jobs:

            title = (
                jd.get(
                    "title"
                )
                if isinstance(
                    jd,
                    dict
                )
                else None
            )

            if (
                not isinstance(
                    title,
                    str
                )
                or not title.strip()
                or len(
                    title.strip()
                ) > 200
                or not scraper.is_job(
                    title
                )
            ):

                invalid_count += 1

                continue

            existing = find_duplicate_job(
                db,
                jd
            )

            if not existing:

                (
                    application_url,
                    recruiter_email
                ) = (
                    validated_scrape_application_fields(
                        jd
                    )
                )

                row = Job(
                    title=title.strip(),

                    company=jd.get(
                        "company"
                    ),

                    job_type=jd.get(
                        "job_type"
                    ),

                    location=jd.get(
                        "location"
                    ),

                    salary=jd.get(
                        "salary"
                    ),

                    description=jd.get(
                        "description"
                    ),

                    source=jd.get(
                        "source",
                        "Custom"
                    ),

                    source_url=(
                        jd.get(
                            "source_url"
                        )
                        or jd.get(
                            "url"
                        )
                        or url
                    ),

                    scraped_at=(
                        parse_scraped_datetime(
                            jd.get(
                                "scraped_at"
                            )
                        )
                        or datetime.utcnow()
                    ),

                    application_url=application_url,

                    recruiter_email=recruiter_email,

                    requirements=jd.get(
                        "requirements"
                    ),

                    responsibilities=jd.get(
                        "responsibilities"
                    ),

                    qualifications=jd.get(
                        "qualifications"
                    ),

                    required_skills=jd.get(
                        "required_skills"
                    ),

                    experience_requirements=jd.get(
                        "experience_requirements"
                    ),

                    application_deadline=jd.get(
                        "application_deadline"
                    ),

                    application_instructions=jd.get(
                        "application_instructions"
                    ),

                    posted_date=(
                        parse_scraped_datetime(
                            jd.get(
                                "posted_date"
                            )
                        )
                    ),
                )

                row.completeness_score = sum(
                    bool(value)
                    for value in [
                        row.title,
                        row.company,
                        row.location,
                        row.job_type,
                        row.salary,
                        row.description,
                        row.requirements,
                        row.responsibilities,
                        row.qualifications,
                        row.required_skills,
                        row.application_url,
                    ]
                )

                db.add(
                    row
                )

                new_count += 1

                linked_count += bool(
                    row.application_url
                    or row.recruiter_email
                )

            else:

                duplicate_count += 1

        if progress_callback:

            progress_callback(
                "Saving validated listings"
            )

        db.commit()

        return {
            "discovered": len(jobs),
            "added": new_count,
            "duplicates": duplicate_count,
            "invalid": invalid_count,
            "with_application_links": linked_count,
            "without_application_links": max(
                0,
                new_count - linked_count
            ),
        }

    finally:

        db.close()


@app.post("/api/scrape-url")
def scrape_custom_url(
    data: URLScrape,
    background_tasks: BackgroundTasks,
    _admin: bool = Depends(require_admin)
):

    try:

        url = scraper.validate_public_http_url(
            data.url
        )

    except ValueError as e:

        raise HTTPException(
            status_code=422,
            detail=str(e)
        ) from e

    if not scrape_lock.acquire(
        blocking=False
    ):

        raise HTTPException(
            status_code=409,
            detail=(
                "A scraping request is already running."
            )
        )

    task_id = uuid.uuid4().hex

    scrape_runs[
        task_id
    ] = {
        "state": "running",
        "url": url,
        "progress": "Fetching listings",
    }

    background_tasks.add_task(
        run_scrape_task,
        task_id,
        url
    )

    return {
        "success": True,
        "task_id": task_id,
        "state": "running",
        "url": url,
    }


# ============================================================
# CV
# ============================================================

CV_TEMPLATES = [
    {
        "id": "modern",
        "name": "Modern Professional",
        "description": (
            "Contemporary sections with one restrained accent color."
        ),
    },
    {
        "id": "ats",
        "name": "ATS Professional",
        "description": (
            "Single-column, text-first layout with no photo by default."
        ),
    },
    {
        "id": "executive",
        "name": "Executive",
        "description": (
            "Concise leadership layout with clear career progression."
        ),
    },
    {
        "id": "creative",
        "name": "Creative Professional",
        "description": (
            "A distinctive hierarchy that remains readable and print-friendly."
        ),
    },
    {
        "id": "minimal",
        "name": "Minimal",
        "description": (
            "Compact typography and simple section dividers."
        ),
    },
]


@app.get("/api/cv/templates")
def get_cv_templates():

    return {
        "success": True,
        "templates": CV_TEMPLATES,
    }


@app.get("/api/cv/status")
def cv_ai_status():

    return {
        "success": True,
        "ai_available": cv_ai_configured(),
        "provider": (
            "openai"
            if cv_ai_configured()
            else None
        ),
    }


def run_cv_ai(
    mode: str,
    request: CVAIRequest
):

    if not cv_ai_configured():

        raise HTTPException(
            status_code=503,
            detail=(
                "CV AI is not configured. "
                "Use the profile-based builder "
                "or configure the OpenAI provider "
                "on the server."
            )
        )

    try:

        updated = update_cv_profile(
            mode,
            request
        )

    except CVAIError as error:

        raise HTTPException(
            status_code=502,
            detail=str(error)
        ) from error

    original = request.profile

    for key in (
        "full_name",
        "email",
        "phone",
        "location",
        "website",
        "linkedin",
        "github",
    ):

        if key in original:

            updated[key] = (
                original[key]
            )

    return {
        "success": True,
        "profile": updated,
        "message": (
            "AI draft ready. "
            "Review every detail before using it."
        ),
    }


@app.post("/api/cv/generate")
def generate_cv(
    request: CVAIRequest
):

    return run_cv_ai(
        "generate",
        request
    )


@app.post("/api/cv/tailor")
def tailor_cv(
    request: CVAIRequest
):

    return run_cv_ai(
        "tailor",
        request
    )


@app.post("/api/cv/instruct")
def instruct_cv(
    request: CVAIRequest
):

    return run_cv_ai(
        "instruct",
        request
    )


def cv_lines(
    value
):

    if isinstance(
        value,
        list
    ):

        return [
            str(item).strip()
            for item in value
            if str(item).strip()
        ]

    if isinstance(
        value,
        str
    ):

        return [
            line.strip()
            for line in value.splitlines()
            if line.strip()
        ]

    return []


@app.post("/api/cv/preview")
def preview_cv(
    request: CVExportRequest
):

    if request.template not in {
        item["id"]
        for item in CV_TEMPLATES
    }:

        raise HTTPException(
            status_code=422,
            detail=(
                "Choose one of the available "
                "CV templates."
            )
        )

    profile = request.profile

    sections = [
        {
            "label": label,
            "items": cv_lines(
                profile.get(key)
            ),
        }

        for key, label in [
            (
                "summary",
                "Professional summary"
            ),
            (
                "experience",
                "Work experience"
            ),
            (
                "internships",
                "Internships"
            ),
            (
                "projects",
                "Projects"
            ),
            (
                "education",
                "Education"
            ),
            (
                "skills",
                "Skills"
            ),
            (
                "certifications",
                "Certifications"
            ),
            (
                "languages",
                "Languages"
            ),
            (
                "achievements",
                "Achievements"
            ),
            (
                "references",
                "References"
            ),
        ]

        if profile.get(key)
    ]

    return {
        "success": True,
        "template": request.template,
        "sections": sections,
        "target_job": request.target_job,
    }


def cv_photo_bytes(
    request
):

    if (
        not request.include_photo
        or not request.photo_base64
    ):

        return None

    encoded = (
        request.photo_base64
        .split(
            ",",
            1
        )[-1]
    )

    if len(encoded) > 4_200_000:

        raise HTTPException(
            status_code=413,
            detail=(
                "Profile photos must be 3 MB or smaller."
            )
        )

    try:

        raw = base64.b64decode(
            encoded,
            validate=True
        )

        from PIL import Image

        image = Image.open(
            BytesIO(raw)
        )

        if (
            image.format
            not in {
                "JPEG",
                "PNG",
            }
            or image.width > 4000
            or image.height > 4000
            or image.width < 80
            or image.height < 80
        ):

            raise ValueError(
                "unsupported image"
            )

        image.verify()

        return raw

    except Exception as error:

        raise HTTPException(
            status_code=415,
            detail=(
                "Use a valid JPG or PNG image "
                "no larger than 4000 by 4000 pixels."
            )
        ) from error


def export_cv_docx(
    request,
    photo
):

    from docx import Document
    from docx.shared import (
        Inches,
        Pt,
        RGBColor
    )
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(
        0.58
        if request.template == "ats"
        else 0.68
    )

    section.bottom_margin = Inches(
        0.62
    )

    section.left_margin = Inches(
        0.7
    )

    section.right_margin = Inches(
        0.7
    )

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(
        9.5
    )

    styles["Normal"].font.color.rgb = (
        RGBColor(
            43,
            58,
            50
        )
    )

    accent = {
        "creative": "4A6091",
        "executive": "40594D",
        "minimal": "414B45",
        "modern": "2D7654",
        "ats": "2D7654",
    }[
        request.template
    ]

    name = str(
        request.profile.get(
            "full_name",
            ""
        )
    ).strip()

    title = str(
        request.profile.get(
            "professional_title",
            ""
        )
    ).strip()

    if photo:

        try:

            document.add_picture(
                BytesIO(photo),
                width=Inches(0.95)
            )

        except Exception:

            pass

    if name:

        p = document.add_paragraph()

        run = p.add_run(
            name
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(
            25
        )

        run.font.color.rgb = (
            RGBColor.from_string(
                accent
            )
        )

    if title:

        p = document.add_paragraph(
            title
        )

        p.paragraph_format.space_after = Pt(
            4
        )

    contact = "  |  ".join(
        str(
            request.profile.get(
                key,
                ""
            )
        ).strip()

        for key in (
            "email",
            "phone",
            "location",
            "website",
            "linkedin",
            "github",
        )

        if str(
            request.profile.get(
                key,
                ""
            )
        ).strip()
    )

    if contact:

        document.add_paragraph(
            contact
        )

    if request.target_job:

        target = " · ".join(
            value
            for value in (
                request.target_job.get(
                    "title"
                ),
                request.target_job.get(
                    "company"
                ),
            )
            if value
        )

        if target:

            p = document.add_paragraph()

            run = p.add_run(
                f"Target role: {target}"
            )

            run.italic = True
            run.font.size = Pt(
                9
            )

    for key, label in [
        (
            "summary",
            "Professional Summary"
        ),
        (
            "experience",
            "Experience"
        ),
        (
            "internships",
            "Internships"
        ),
        (
            "projects",
            "Projects"
        ),
        (
            "education",
            "Education"
        ),
        (
            "skills",
            "Skills"
        ),
        (
            "certifications",
            "Certifications"
        ),
        (
            "languages",
            "Languages"
        ),
        (
            "achievements",
            "Achievements"
        ),
        (
            "references",
            "References"
        ),
    ]:

        items = cv_lines(
            request.profile.get(
                key
            )
        )

        if not items:
            continue

        heading = document.add_paragraph()

        heading.paragraph_format.space_before = Pt(
            9
        )

        heading.paragraph_format.space_after = Pt(
            3
        )

        run = heading.add_run(
            label.upper()
        )

        run.bold = True
        run.font.size = Pt(
            10
        )

        run.font.color.rgb = (
            RGBColor.from_string(
                accent
            )
        )

        for item in items:

            paragraph = document.add_paragraph(
                style="Normal"
            )

            paragraph.paragraph_format.space_after = Pt(
                3
            )

            paragraph.paragraph_format.keep_together = True

            paragraph.add_run(
                item
            )

    footer = section.footer.paragraphs[0]

    footer.alignment = 2

    run = footer.add_run(
        "Grafan CV · Page "
    )

    field = OxmlElement(
        "w:fldSimple"
    )

    field.set(
        qn("w:instr"),
        "PAGE"
    )

    run._r.addnext(
        field
    )

    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


def export_cv_pdf(
    request,
    photo
):

    try:

        import fitz

    except ImportError as error:

        raise HTTPException(
            status_code=503,
            detail=(
                "PDF export dependency is "
                "not installed on this server."
            )
        ) from error

    document = fitz.open()

    page = document.new_page(
        width=595,
        height=842
    )

    accent = {
        "creative": (
            0.29,
            0.38,
            0.57
        ),

        "executive": (
            0.25,
            0.35,
            0.30
        ),

        "minimal": (
            0.25,
            0.29,
            0.27
        ),

        "modern": (
            0.18,
            0.46,
            0.33
        ),

        "ats": (
            0.18,
            0.46,
            0.33
        ),
    }[
        request.template
    ]

    x = 52
    y = 58
    bottom = 785

    font = "helv"

    def add_page():

        nonlocal page, y

        page = document.new_page(
            width=595,
            height=842
        )

        y = 55

    def write(
        text,
        size=10,
        color=(
            0.17,
            0.23,
            0.20
        ),
        bold=False,
        gap=3
    ):

        nonlocal y

        width = 495

        words = str(text).split()

        lines = []
        line = ""

        for word in words:

            candidate = (
                f"{line} {word}"
            ).strip()

            if fitz.get_text_length(
                candidate,
                fontname=(
                    "hebo"
                    if bold
                    else font
                ),
                fontsize=size
            ) <= width:

                line = candidate

            else:

                if line:

                    lines.append(
                        line
                    )

                line = word

        if line:

            lines.append(
                line
            )

        for row in (
            lines
            or [""]
        ):

            if (
                y
                + size
                + gap
                > bottom
            ):

                add_page()

            page.insert_text(
                (
                    x,
                    y
                ),
                row,
                fontname=(
                    "hebo"
                    if bold
                    else font
                ),
                fontsize=size,
                color=color,
            )

            y += (
                size
                * 1.45
            )

        y += gap

    profile = request.profile

    if photo:

        page.insert_image(
            fitz.Rect(
                475,
                45,
                540,
                110
            ),
            stream=photo,
            keep_proportion=True,
        )

    write(
        profile.get(
            "full_name",
            ""
        ),
        24,
        accent,
        True,
        2
    )

    write(
        profile.get(
            "professional_title",
            ""
        ),
        12,
        (
            0.3,
            0.36,
            0.33
        ),
        False,
        3
    )

    contact = " | ".join(
        str(
            profile.get(
                key,
                ""
            )
        ).strip()

        for key in (
            "email",
            "phone",
            "location",
            "website",
            "linkedin",
            "github",
        )

        if str(
            profile.get(
                key,
                ""
            )
        ).strip()
    )

    if contact:

        write(
            contact,
            8,
            (
                0.36,
                0.42,
                0.38
            ),
            False,
            7
        )

    if request.target_job:

        target = " · ".join(
            value
            for value in (
                request.target_job.get(
                    "title"
                ),
                request.target_job.get(
                    "company"
                ),
            )
            if value
        )

        if target:

            write(
                f"Target role: {target}",
                9,
                (
                    0.32,
                    0.37,
                    0.34
                ),
                False,
                5
            )

    for key, label in [
        (
            "summary",
            "Professional Summary"
        ),
        (
            "experience",
            "Experience"
        ),
        (
            "internships",
            "Internships"
        ),
        (
            "projects",
            "Projects"
        ),
        (
            "education",
            "Education"
        ),
        (
            "skills",
            "Skills"
        ),
        (
            "certifications",
            "Certifications"
        ),
        (
            "languages",
            "Languages"
        ),
        (
            "achievements",
            "Achievements"
        ),
        (
            "references",
            "References"
        ),
    ]:

        items = cv_lines(
            profile.get(
                key
            )
        )

        if not items:

            continue

        write(
            label.upper(),
            10,
            accent,
            True,
            3
        )

        for item in items:

            write(
                item,
                9.5,
                (
                    0.17,
                    0.23,
                    0.20
                ),
                False,
                4
            )

    for number, item in enumerate(
        document,
        start=1
    ):

        item.insert_text(
            (
                500,
                815
            ),
            str(number),
            fontname="helv",
            fontsize=8,
            color=(
                0.45,
                0.49,
                0.46
            ),
        )

    return BytesIO(
        document.tobytes()
    )


@app.post("/api/cv/export")
def export_cv(
    request: CVExportRequest
):

    if request.template not in {
        item["id"]
        for item in CV_TEMPLATES
    }:

        raise HTTPException(
            status_code=422,
            detail=(
                "Choose one of the available "
                "CV templates."
            )
        )

    if not str(
        request.profile.get(
            "full_name",
            ""
        )
    ).strip():

        raise HTTPException(
            status_code=422,
            detail=(
                "Add your name before exporting your CV."
            )
        )

    if sum(
        len(value)
        for value in request.profile.values()
        if isinstance(
            value,
            str
        )
    ) > 60000:

        raise HTTPException(
            status_code=413,
            detail=(
                "CV information exceeds "
                "the 60,000 character limit."
            )
        )

    photo = cv_photo_bytes(
        request
    )

    if (
        request.format.casefold()
        == "docx"
    ):

        output = export_cv_docx(
            request,
            photo
        )

        media_type = (
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )

        extension = "docx"

    elif (
        request.format.casefold()
        == "pdf"
    ):

        output = export_cv_pdf(
            request,
            photo
        )

        media_type = (
            "application/pdf"
        )

        extension = "pdf"

    else:

        raise HTTPException(
            status_code=422,
            detail="Choose PDF or DOCX format."
        )

    return StreamingResponse(
        output,
        media_type=media_type,
        headers={
            "Content-Disposition": (
                f'attachment; filename="grafan-cv.{extension}"'
            ),
            "Cache-Control": "no-store",
        },
    )


@app.post("/api/cv/parse")
@app.post("/api/cv/upload")
@app.post("/api/cv/analyze")
async def parse_cv(
    file: UploadFile = File(...)
):

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="No file uploaded"
        )

    extension = Path(
        file.filename
    ).suffix.casefold()

    supported_types = {
        ".pdf": {
            "application/pdf",
            "application/octet-stream",
        },

        ".docx": {
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document",

            "application/octet-stream",
        },
    }

    if (
        extension not in supported_types
        or file.content_type
        not in supported_types[
            extension
        ]
    ):

        raise HTTPException(
            status_code=415,
            detail="Upload a PDF or DOCX file."
        )

    content = await file.read(
        MAX_UPLOAD_BYTES + 1
    )

    if (
        len(content)
        > MAX_UPLOAD_BYTES
    ):

        raise HTTPException(
            status_code=413,
            detail=(
                "CV files must be "
                "10 MB or smaller."
            )
        )

    if (
        extension == ".pdf"
        and not content.startswith(
            b"%PDF-"
        )
    ):

        raise HTTPException(
            status_code=415,
            detail=(
                "The uploaded file is not "
                "a valid PDF."
            )
        )

    if (
        extension == ".docx"
        and not content.startswith(
            b"PK"
        )
    ):

        raise HTTPException(
            status_code=415,
            detail=(
                "The uploaded file is not "
                "a valid DOCX document."
            )
        )

    filepath = os.path.join(
        UPLOAD_FOLDER,
        f"{uuid.uuid4().hex}{extension}"
    )

    try:

        with open(
            filepath,
            "wb"
        ) as f:

            f.write(
                content
            )

        text = (
            cv_parser.parse_pdf(
                filepath
            )
            if extension == ".pdf"
            else cv_parser.parse_docx(
                filepath
            )
        )

        if not text.strip():

            raise HTTPException(
                status_code=422,
                detail=(
                    "No selectable text was found. "
                    "Scanned documents are not supported "
                    "for CV analysis."
                )
            )

        info = cv_parser.extract_info(
            text
        )

        return {
            "success": True,
            "extracted": info,
        }

    except HTTPException:

        raise

    except Exception as e:

        print(
            f"CV analysis failed "
            f"({type(e).__name__})."
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "The CV could not be analyzed. "
                "Check the file and try again."
            )
        ) from e

    finally:

        if os.path.exists(
            filepath
        ):

            os.remove(
                filepath
            )


# ============================================================
# APPLY - FIXED
# ============================================================

@app.post("/api/apply")
async def apply_job(
    job_id: int = Form(...),
    applicant_name: str = Form(...),
    applicant_email: str = Form(...),
    applicant_phone: str = Form(""),
    cover_letter: str = Form(""),
    cv: Optional[UploadFile] = File(None),
):

    """
    Submit an application through Grafan.

    Workflow:

    1. Validate the selected job.
    2. Validate applicant information.
    3. Save the applicant CV.
    4. Save the application in SQLite.
    5. Send application + CV to recruiter.
    6. Send confirmation email to applicant.
    7. Return employer application URL when available.
    """

    db = SessionLocal()

    cv_path = None

    try:

        # ------------------------------------------------------
        # FIND JOB
        # ------------------------------------------------------

        job = db.get(
            Job,
            job_id
        )

        if (
            not job
            or not is_displayable_job(job)
        ):

            raise HTTPException(
                status_code=404,
                detail=(
                    "This job is no longer available."
                )
            )

        # ------------------------------------------------------
        # CLEAN APPLICANT INPUT
        # ------------------------------------------------------

        applicant_name = (
            applicant_name
            or ""
        ).strip()

        applicant_email = (
            applicant_email
            or ""
        ).strip().lower()

        applicant_phone = (
            applicant_phone
            or ""
        ).strip()

        cover_letter = (
            cover_letter
            or ""
        ).strip()

        # ------------------------------------------------------
        # VALIDATE NAME
        # ------------------------------------------------------

        if not applicant_name:

            raise HTTPException(
                status_code=422,
                detail="Full name is required."
            )

        if len(
            applicant_name
        ) > 100:

            raise HTTPException(
                status_code=422,
                detail="Full name is too long."
            )

        # ------------------------------------------------------
        # VALIDATE EMAIL
        # ------------------------------------------------------

        if not re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            applicant_email
        ):

            raise HTTPException(
                status_code=422,
                detail=(
                    "Please provide a valid "
                    "email address."
                )
            )

        # ------------------------------------------------------
        # VALIDATE PHONE
        # ------------------------------------------------------

        if len(
            applicant_phone
        ) > 20:

            raise HTTPException(
                status_code=422,
                detail="Phone number is too long."
            )

        # ------------------------------------------------------
        # VALIDATE COVER LETTER
        # ------------------------------------------------------

        if len(
            cover_letter
        ) > 12000:

            raise HTTPException(
                status_code=422,
                detail="Cover letter is too long."
            )

        # ------------------------------------------------------
        # SAVE CV
        # ------------------------------------------------------

        cv_filename = None

        if cv and cv.filename:

            original_name = Path(
                cv.filename
            ).name

            extension = Path(
                original_name
            ).suffix.casefold()

            allowed_extensions = {
                ".pdf",
                ".doc",
                ".docx",
            }

            if extension not in allowed_extensions:

                raise HTTPException(
                    status_code=400,
                    detail=(
                        "CV must be PDF, DOC, or DOCX."
                    )
                )

            cv_bytes = await cv.read(
                MAX_UPLOAD_BYTES + 1
            )

            if (
                len(cv_bytes)
                > MAX_UPLOAD_BYTES
            ):

                raise HTTPException(
                    status_code=413,
                    detail=(
                        "CV file must be "
                        "10 MB or smaller."
                    )
                )

            safe_name = re.sub(
                r"[^A-Za-z0-9._-]+",
                "_",
                original_name
            ).strip("._")

            if not safe_name:

                safe_name = (
                    f"cv{extension}"
                )

            stored_name = (
                f"{uuid.uuid4().hex}_"
                f"{safe_name}"
            )

            upload_dir = Path(
                UPLOAD_FOLDER
            )

            upload_dir.mkdir(
                parents=True,
                exist_ok=True
            )

            cv_path = (
                upload_dir
                / stored_name
            )

            with open(
                cv_path,
                "wb"
            ) as uploaded_file:

                uploaded_file.write(
                    cv_bytes
                )

            cv_filename = (
                original_name[:200]
            )

        # ------------------------------------------------------
        # SAVE APPLICATION
        # ------------------------------------------------------

        application = Application(

            job_id=job.id,

            applicant_name=(
                applicant_name
            ),

            applicant_email=(
                applicant_email
            ),

            applicant_phone=(
                applicant_phone
            ),

            cv_filename=(
                cv_filename
            ),

            cover_letter=(
                cover_letter
            ),

            applied_date=datetime.utcnow(),
        )

        db.add(
            application
        )

        db.commit()

        db.refresh(
            application
        )

        # ------------------------------------------------------
        # FIND RECRUITER EMAIL
        # ------------------------------------------------------

        recruiter_email = None

        raw_recruiter_email = (
            job.recruiter_email
            or ""
        ).strip()

        if re.fullmatch(
            r"[^\s@]+@[^\s@]+\.[^\s@]+",
            raw_recruiter_email
        ):

            recruiter_email = (
                raw_recruiter_email
            )

        # ------------------------------------------------------
        # EMAIL STATUS
        # ------------------------------------------------------

        email_enabled = (
            os.environ.get(
                "EMAIL_NOTIFICATIONS",
                "false"
            ).strip().lower()
            in {
                "1",
                "true",
                "yes",
                "on",
            }
            and email_sender.is_configured
        )

        recruiter_notified = False

        confirmation_sent = False

        # ------------------------------------------------------
        # SEND EMAILS
        # ------------------------------------------------------

        if email_enabled:

            # --------------------------------------------------
            # EMAIL RECRUITER
            # --------------------------------------------------

            if recruiter_email:

                recruiter_notified = (
                    await asyncio.to_thread(
                        email_sender.send_application_email,

                        recruiter_email,

                        job.title
                        or "",

                        job.company
                        or "",

                        applicant_name,

                        applicant_email,

                        applicant_phone,

                        cover_letter,

                        (
                            str(cv_path)
                            if cv_path
                            else None
                        ),
                    )
                )

            # --------------------------------------------------
            # EMAIL APPLICANT CONFIRMATION
            # --------------------------------------------------

            confirmation_sent = (
                await asyncio.to_thread(
                    email_sender.send_confirmation_email,

                    applicant_email,

                    applicant_name,

                    job.title
                    or "",

                    job.company
                    or "",
                )
            )

        # ------------------------------------------------------
        # ORIGINAL EMPLOYER URL
        # ------------------------------------------------------

        application_url = None

        if job.application_url:

            try:

                parsed = urlparse(
                    job.application_url.strip()
                )

                if (
                    parsed.scheme
                    in {
                        "http",
                        "https",
                    }
                    and parsed.netloc
                ):

                    application_url = (
                        job.application_url.strip()
                    )

            except Exception:

                application_url = None

        # ------------------------------------------------------
        # RESPONSE
        # ------------------------------------------------------

        return {
            "success": True,

            "message": (
                "Application submitted successfully."
            ),

            "application_id": (
                application.id
            ),

            "recruiter_notified": (
                recruiter_notified
            ),

            "confirmation_sent": (
                confirmation_sent
            ),

            "application_url": (
                application_url
            ),
        }

    except HTTPException:

        db.rollback()

        raise

    except Exception as error:

        db.rollback()

        print(
            "Application submission failed "
            f"({type(error).__name__}): "
            f"{error}"
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "The application could not be submitted."
            )
        ) from error

    finally:

        db.close()


# ============================================================
# STATS
# ============================================================

@app.get("/api/stats")
def get_stats():

    db = SessionLocal()

    try:

        expire_due_ads(
            db
        )

        jobs = [
            job
            for job in db.query(
                Job
            ).all()
            if is_displayable_job(job)
        ]

        active_ads = [
            ad
            for ad in (
                db.query(
                    Advertisement
                )
                .filter_by(
                    status="published"
                )
                .all()
            )
            if not ad_is_expired(ad)
        ]

        return {

            "total_jobs": len(
                jobs
            ),

            "companies": len(
                {
                    job.company.strip().casefold()
                    for job in jobs
                    if (
                        job.company
                        and job.company.strip()
                    )
                }
            ),

            "advertisements": len(
                active_ads
            ),

            "jobs_by_type": {

                "remote": sum(
                    job.job_type == "remote"
                    for job in jobs
                ),

                "hybrid": sum(
                    job.job_type == "hybrid"
                    for job in jobs
                ),

                "onsite": sum(
                    job.job_type == "onsite"
                    for job in jobs
                ),

                "gig": sum(
                    job.job_type == "gig"
                    for job in jobs
                ),
            },
        }

    finally:

        db.close()


# ============================================================
# SOURCES
# ============================================================

@app.get("/api/sources")
def get_sources(
    _admin: bool = Depends(require_admin)
):

    return {
        "success": True,
        "sources": (
            scraper.get_configured_sources()
        ),
    }


@app.get("/api/admin/jobs")
def admin_jobs(
    search: str = "",
    source: str = "",
    status: str = "all",
    limit: int = 100,
    offset: int = 0,
    _admin: bool = Depends(require_admin),
):

    if (
        limit < 1
        or limit > 500
        or offset < 0
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Invalid admin job-list "
                "page size or offset."
            )
        )

    db = SessionLocal()

    try:

        query = db.query(
            Job
        )

        if search.strip():

            term = (
                f"%{search.strip()}%"
            )

            query = query.filter(
                Job.title.ilike(term)
                | Job.company.ilike(term)
            )

        if source.strip():

            query = query.filter(
                Job.source == source.strip()
            )

        if status in {
            "active",
            "inactive",
        }:

            query = query.filter(
                Job.job_status == status
            )

        elif status != "all":

            raise HTTPException(
                status_code=422,
                detail=(
                    "Status must be "
                    "active, inactive, or all."
                )
            )

        total = query.count()

        jobs = (
            query
            .order_by(
                Job.posted_date.desc(),
                Job.id.desc()
            )
            .offset(offset)
            .limit(limit)
            .all()
        )

        identities = {}

        for row in (
            db.query(
                Job.id,
                Job.title,
                Job.company
            )
            .order_by(
                Job.id.asc()
            )
            .all()
        ):

            key = (
                re.sub(
                    r"\W+",
                    " ",
                    (
                        row.title
                        or ""
                    ).casefold()
                ).strip(),

                re.sub(
                    r"\W+",
                    " ",
                    (
                        row.company
                        or ""
                    ).casefold()
                ).strip(),
            )

            identities.setdefault(
                key,
                row.id
            )

        return {
            "success": True,
            "total": total,

            "jobs": [
                {
                    **admin_job_payload(
                        job
                    ),

                    "duplicate_status": (
                        "duplicate"

                        if identities.get(
                            (
                                re.sub(
                                    r"\W+",
                                    " ",
                                    (
                                        job.title
                                        or ""
                                    ).casefold()
                                ).strip(),

                                re.sub(
                                    r"\W+",
                                    " ",
                                    (
                                        job.company
                                        or ""
                                    ).casefold()
                                ).strip(),
                            )
                        )
                        != job.id

                        else "unique"
                    ),
                }

                for job in jobs
            ],
        }

    finally:

        db.close()


@app.get("/api/admin/sources")
def admin_sources(
    _admin: bool = Depends(require_admin)
):

    return {
        "success": True,
        "sources": (
            scraper.get_configured_sources()
        ),
    }


@app.get("/api/admin/scraping/stats")
def admin_scraping_stats(
    _admin: bool = Depends(require_admin)
):

    db = SessionLocal()

    try:

        jobs = (
            db.query(
                Job
            )
            .all()
        )

        groups = {}

        for job in jobs:

            key = (
                re.sub(
                    r"\W+",
                    " ",
                    (
                        job.title
                        or ""
                    ).casefold()
                ).strip(),

                re.sub(
                    r"\W+",
                    " ",
                    (
                        job.company
                        or ""
                    ).casefold()
                ).strip(),
            )

            groups[key] = (
                groups.get(
                    key,
                    0
                )
                + 1
            )

        return {

            "success": True,

            "total_jobs": len(
                jobs
            ),

            "active_jobs": sum(
                job.job_status == "active"
                for job in jobs
            ),

            "scraped_today": sum(
                bool(
                    job.scraped_at
                    and
                    job.scraped_at.date()
                    == datetime.now().date()
                )
                for job in jobs
            ),

            "with_application_links": sum(
                bool(
                    job.application_url
                    or job.recruiter_email
                )
                for job in jobs
            ),

            "without_application_links": sum(
                not (
                    job.application_url
                    or job.recruiter_email
                )
                for job in jobs
            ),

            "duplicates": sum(
                count - 1
                for count in groups.values()
                if count > 1
            ),

            "jobs_by_source": {
                source or "Unknown": count
                for source, count in (
                    db.query(
                        Job.source,
                        func.count(Job.id)
                    )
                    .group_by(
                        Job.source
                    )
                    .all()
                )
            },
        }

    finally:

        db.close()


def create_source_scrape_task(
    source_name: str,
    background_tasks: BackgroundTasks
):

    from scraper import CONFIGURED_SOURCES

    configured = {
        name: method
        for name, _url, method
        in CONFIGURED_SOURCES
    }

    method_name = configured.get(
        source_name
    )

    if not method_name:

        raise HTTPException(
            status_code=404,
            detail=(
                "That source is not configured."
            )
        )

    if not scrape_lock.acquire(
        blocking=False
    ):

        raise HTTPException(
            status_code=409,
            detail=(
                "A scraping request is already running."
            )
        )

    task_id = uuid.uuid4().hex

    scrape_runs[
        task_id
    ] = {
        "state": "running",
        "source": source_name,
        "progress": "Fetching listings",
    }

    background_tasks.add_task(
        run_source_scrape_task,
        task_id,
        source_name,
        method_name
    )

    return {
        "success": True,
        "task_id": task_id,
        "state": "running",
    }


@app.post("/api/admin/scrape/all")
def admin_scrape_all(
    background_tasks: BackgroundTasks,
    _admin: bool = Depends(require_admin)
):

    return trigger_scrape(
        background_tasks,
        True
    )


@app.post("/api/admin/scrape/source")
def admin_scrape_source(
    data: AdminScrapeSource,
    background_tasks: BackgroundTasks,
    _admin: bool = Depends(require_admin)
):

    return create_source_scrape_task(
        data.source,
        background_tasks
    )


@app.post("/api/admin/scrape/url")
def admin_scrape_url(
    data: AdminScrapeURL,
    background_tasks: BackgroundTasks,
    _admin: bool = Depends(require_admin)
):

    try:

        url = scraper.validate_public_http_url(
            data.url
        )

    except ValueError as error:

        raise HTTPException(
            status_code=422,
            detail=str(error)
        ) from error

    return trigger_scrape(
        background_tasks,
        True,
        url
    )


@app.get("/api/admin/scrape/status/{task_id}")
def admin_scrape_status(
    task_id: str,
    _admin: bool = Depends(require_admin)
):

    return get_scrape_status(
        task_id,
        True
    )


@app.post("/api/admin/jobs", status_code=201)
def admin_create_job(
    job: JobCreate,
    _admin: bool = Depends(require_admin)
):

    return create_job(
        job,
        True
    )


@app.put("/api/admin/jobs/{job_id}")
def admin_update_job(
    job_id: int,
    data: JobUpdate,
    _admin: bool = Depends(require_admin)
):

    return update_job(
        job_id,
        data,
        True
    )


@app.delete("/api/admin/jobs/{job_id}")
def admin_delete_job(
    job_id: int,
    _admin: bool = Depends(require_admin)
):

    return delete_job(
        job_id,
        True
    )


# ============================================================
# AI CAREER COACH
# ============================================================

CAREER_TOPICS = {

    "python": (
        "Software and data roles",

        "Build a small portfolio that shows how you use Python "
        "to solve practical problems. Choose projects related "
        "to the roles you want, and explain your decisions and results.",

        [
            "Review current job descriptions for the Python tools they request.",
            "Publish one or two finished projects with clear readme files.",
            "Practice explaining a project from problem to outcome.",
        ],
    ),

    "java": (
        "Java roles",

        "Focus your applications on the kind of Java work you want, "
        "such as backend services or Android development. Show how "
        "your code is tested and maintained.",

        [
            "Review current descriptions for frameworks and tools employers request.",
            "Build a small project with tests and clear setup instructions.",
            "Prepare examples of debugging and teamwork.",
        ],
    ),

    "javascript": (
        "Web development",

        "Show a working web project that demonstrates how you build "
        "accessible, responsive experiences and connect them to data.",

        [
            "Choose a project related to the roles you want.",
            "Document the tools you used and the decisions you made.",
            "Practice explaining how you test and improve the project.",
        ],
    ),

    "data": (
        "Data roles",

        "Employers need people who can turn data into clear decisions. "
        "Make your analysis process and conclusions easy to follow.",

        [
            "Practice SQL and spreadsheet analysis on a relevant dataset.",
            "Create a short report or dashboard with a clear takeaway.",
            "Explain how you checked the data and reached your conclusion.",
        ],
    ),

    "cv": (
        "CV improvement",

        "Make your CV easy to scan and tailor it to each position. "
        "Describe outcomes and evidence instead of listing responsibilities alone.",

        [
            "Put your contact details and a short professional summary near the top.",
            "Use action verbs and concrete results in your experience section.",
            "Mirror relevant terms from the job description when they accurately describe your experience.",
        ],
    ),

    "interview": (
        "Interview preparation",

        "Prepare concise examples that show how you approach a challenge, "
        "work with others, and learn from results.",

        [
            "Choose three examples from your work, studies, or volunteering.",
            "Structure each answer around the situation, your action, and the result.",
            "Prepare thoughtful questions about the role and team.",
        ],
    ),

    "salary": (
        "Salary conversations",

        "Pay varies by role, experience, employer, and location. "
        "Use current comparable listings and your responsibilities "
        "to prepare a realistic range.",

        [
            "Research several recent listings for similar work.",
            "List your relevant experience and the value you bring.",
            "Ask about the full package and responsibilities before deciding.",
        ],
    ),

    "remote": (
        "Remote work",

        "Remote roles often require clear communication, reliable planning, "
        "and comfort working independently across a distributed team.",

        [
            "Show examples of how you organize work and communicate progress.",
            "Check the role’s location, time-zone, and connectivity requirements.",
            "Prepare questions about onboarding and team communication.",
        ],
    ),

    "accounting": (
        "Accounting and finance",

        "Highlight accuracy, reporting, controls, and the tools you use. "
        "Match your examples to the responsibilities in the role.",

        [
            "List relevant systems, standards, and certifications you actually use.",
            "Prepare an example of a report or process you improved.",
            "Review the employer’s sector and reporting needs.",
        ],
    ),

    "engineering": (
        "Engineering roles",

        "Make your practical experience visible. Explain the constraints, "
        "standards, tools, and outcomes behind your work.",

        [
            "Select a project that demonstrates your technical decisions.",
            "List relevant tools, standards, or professional registration accurately.",
            "Prepare to explain safety, quality, or testing practices you followed.",
        ],
    ),

    "marketing": (
        "Marketing roles",

        "Connect your work to an audience and a measurable outcome. "
        "Include the channel, approach, and evidence where you have it.",

        [
            "Build a short portfolio of campaigns or content you contributed to.",
            "Show the goal, audience, and result for each example.",
            "Review the employer’s audience and products before applying.",
        ],
    ),

    "hr": (
        "Human resources",

        "Show how you support people processes, communication, and fair, "
        "consistent ways of working.",

        [
            "Describe relevant recruitment, people operations, or training work.",
            "Highlight systems and employment practices you know.",
            "Prepare an example that shows confidentiality and good judgment.",
        ],
    ),

    "health": (
        "Healthcare",

        "Make your clinical or public health experience, training, "
        "and professional registration clear for the role.",

        [
            "Check the role’s qualification and registration requirements.",
            "Describe relevant patient, community, or operational experience.",
            "Prepare an example of safe, respectful teamwork.",
        ],
    ),

    "education": (
        "Education",

        "Connect your teaching or education experience to learner needs, "
        "subject knowledge, and outcomes.",

        [
            "Describe the learners and subjects you have supported.",
            "Include relevant qualifications and teaching methods.",
            "Prepare an example of adapting your approach for learners.",
        ],
    ),

    "law": (
        "Legal roles",

        "Highlight your legal training, relevant practice area, research, "
        "and professional requirements for the position.",

        [
            "Check the role’s registration or admission requirements.",
            "Prepare an example of careful research or drafting.",
            "Describe how you manage confidential information.",
        ],
    ),

    "agriculture": (
        "Agriculture",

        "Connect your practical or technical experience to the crop, "
        "livestock, supply chain, or community needs in the role.",

        [
            "Identify the sector and field conditions relevant to the role.",
            "Describe a practical project or measurable result.",
            "List relevant technical training and tools.",
        ],
    ),

    "environment": (
        "Environment and conservation",

        "Show the environmental work you have contributed to and how you "
        "measure or document its outcomes.",

        [
            "Review the employer’s current programmes.",
            "Prepare an example involving research, fieldwork, or community engagement.",
            "Explain the methods and evidence behind your work.",
        ],
    ),

    "social": (
        "Community and social work",

        "Focus on the communities you have supported, the methods you use, "
        "and the outcomes you can evidence.",

        [
            "Describe relevant community or protection work.",
            "Explain how you build trust and handle sensitive information.",
            "Include relevant safeguarding training.",
        ],
    ),

    "management": (
        "Leadership and management",

        "Use specific examples to show how you set priorities, support a team, "
        "and deliver work through others.",

        [
            "Describe the team or project context.",
            "Explain the decisions you made and the result.",
            "Show how you handled feedback or changing priorities.",
        ],
    ),

    "civil": (
        "Civil engineering",

        "Make project responsibilities, technical standards, and site or design experience clear.",

        [
            "Select a relevant project and describe your contribution.",
            "List tools, codes, or registration accurately.",
            "Prepare to discuss safety and quality controls.",
        ],
    ),

    "electrical": (
        "Electrical engineering",

        "Connect your technical experience to the systems, standards, and safety needs in the role.",

        [
            "Describe a relevant design, maintenance, or installation project.",
            "List the tools and standards you used.",
            "Prepare an example focused on testing and safety.",
        ],
    ),

    "mechanical": (
        "Mechanical engineering",

        "Show the equipment, systems, or processes you have worked on and the outcomes of your work.",

        [
            "Describe a relevant project or maintenance task.",
            "List design and analysis tools you use.",
            "Prepare an example involving reliability or safety.",
        ],
    ),

    "default": (
        "Career planning",

        "Choose a clear next step based on the work you want and the evidence "
        "you can build. Use current local job descriptions to guide your plan.",

        [
            "Identify two or three roles that interest you.",
            "Compare their recurring requirements with your current experience.",
            "Choose one practical skill or project to strengthen next.",
        ],
    ),
}


CAREER_KEYWORDS = {

    "python": [
        "python",
        "django",
        "flask",
    ],

    "java": [
        "java",
        "spring",
        "android",
    ],

    "javascript": [
        "javascript",
        "react",
        "vue",
        "angular",
        "node",
    ],

    "data": [
        "data",
        "analytics",
        "sql",
        "tableau",
        "power bi",
    ],

    "cv": [
        "cv",
        "resume",
        "curriculum vitae",
    ],

    "interview": [
        "interview",
        "interviewing",
    ],

    "salary": [
        "salary",
        "pay",
        "compensation",
        "benefits",
    ],

    "remote": [
        "remote",
        "work from home",
        "wfh",
    ],

    "accounting": [
        "accounting",
        "accountant",
        "finance",
        "audit",
        "tax",
    ],

    "engineering": [
        "engineering",
        "engineer",
        "civil",
        "electrical",
        "mechanical",
    ],

    "marketing": [
        "marketing",
        "sales",
        "brand",
    ],

    "hr": [
        "human resource",
        "recruitment",
        "talent",
    ],

    "health": [
        "health",
        "medical",
        "doctor",
        "nurse",
        "clinical",
    ],

    "education": [
        "education",
        "teacher",
        "lecturer",
        "professor",
    ],

    "law": [
        "law",
        "legal",
        "lawyer",
        "advocate",
    ],

    "agriculture": [
        "agriculture",
        "agronomist",
        "farming",
    ],

    "environment": [
        "environment",
        "climate",
        "sustainability",
    ],

    "social": [
        "social work",
        "community development",
    ],

    "management": [
        "management",
        "manager",
        "supervisor",
        "leadership",
    ],

    "civil": [
        "civil",
        "construction",
        "infrastructure",
        "structural",
    ],

    "electrical": [
        "electrical",
        "electronics",
        "power",
    ],

    "mechanical": [
        "mechanical",
        "machinery",
        "hvac",
    ],
}


class RuleBasedCareerCoach:
    """
    Server-side guidance provider;
    the API response remains provider-neutral.
    """

    def advise(
        self,
        question,
        skills=""
    ):

        topic_key = "default"

        for key, terms in CAREER_KEYWORDS.items():

            if any(
                re.search(
                    r"(?<![a-z0-9])"
                    + re.escape(
                        term
                    )
                    + r"(?![a-z0-9])",
                    question
                )
                for term in terms
            ):

                topic_key = key

                break

        topic, summary, next_steps = (
            CAREER_TOPICS[
                topic_key
            ]
        )

        clean_skills = [
            item.strip()
            for item in skills.split(",")
            if item.strip()
        ][:12]

        if clean_skills:

            summary += (
                " Skills you shared: "
                + ", ".join(
                    clean_skills
                )
                + "."
            )

        return {
            "topic": topic,
            "summary": summary,
            "next_steps": next_steps,
        }


career_coach = (
    RuleBasedCareerCoach()
)


@app.post("/api/ai/ask")
def ai_ask(
    data: AIQuestion
):

    question = (
        data.question
        or ""
    ).strip()

    if not question:

        raise HTTPException(
            status_code=400,
            detail=(
                "Write a question "
                "to get career guidance."
            )
        )

    if (
        len(question) > 2000
        or len(
            data.skills
            or ""
        ) > 1000
    ):

        raise HTTPException(
            status_code=422,
            detail=(
                "Your question or "
                "skills list is too long."
            )
        )

    advice = career_coach.advise(
        question.lower(),
        data.skills
        or ""
    )

    return {
        "success": True,
        "response": advice[
            "summary"
        ],
        **advice,
        "provider": "rule_based",
        "timestamp": datetime.now().isoformat(),
    }


# ============================================================
# JOB MATCHER
# ============================================================

KNOWN_SKILLS = {
    "python",
    "javascript",
    "java",
    "react",
    "vue",
    "angular",
    "node",
    "sql",
    "mongodb",
    "aws",
    "docker",
    "kubernetes",
    "management",
    "finance",
    "accounting",
    "hr",
    "marketing",
    "engineering",
    "civil",
    "electrical",
    "mechanical",
    "data",
    "analytics",
    "design",
    "ux",
    "ui",
    "html",
    "css",
    "php",
    "ruby",
    "c++",
    "c#",
    "swift",
    "kotlin",
    "go",
    "rust",
    "typescript",
    "nextjs",
    "leadership",
    "communication",
    "coordination",
    "planning",
    "strategy",
}


FILLER_WORDS = {
    "i",
    "have",
    "skill",
    "skil",
    "skills",
    "in",
    "am",
    "a",
    "an",
    "the",
    "of",
    "for",
    "with",
    "at",
    "from",
    "by",
    "to",
    "and",
    "or",
    "but",
    "on",
    "as",
    "my",
    "your",
    "his",
    "her",
    "our",
    "their",
    "me",
    "you",
    "need",
    "want",
    "looking",
    "job",
    "work",
    "career",
    "experience",
    "knowledge",
    "proficient",
    "expert",
    "ability",
    "capable",
    "familiar",
}


SYNONYMS = {

    "python": [
        "python",
        "py",
        "django",
        "flask",
    ],

    "javascript": [
        "javascript",
        "js",
        "node",
        "react",
        "vue",
        "angular",
    ],

    "java": [
        "java",
        "spring",
    ],

    "react": [
        "react",
        "reactjs",
        "nextjs",
    ],

    "management": [
        "management",
        "manager",
        "supervisor",
        "leadership",
        "coordination",
    ],

    "finance": [
        "finance",
        "financial",
        "accounting",
        "accountant",
        "audit",
    ],

    "developer": [
        "developer",
        "programmer",
        "engineer",
    ],

    "analyst": [
        "analyst",
        "analysis",
        "analytics",
    ],

    "marketing": [
        "marketing",
        "sales",
        "brand",
    ],

    "hr": [
        "hr",
        "human resource",
        "recruitment",
        "talent",
    ],

    "health": [
        "health",
        "medical",
        "doctor",
        "nurse",
        "clinical",
    ],

    "engineering": [
        "engineering",
        "engineer",
        "technical",
    ],

    "civil": [
        "civil",
        "construction",
        "structural",
    ],

    "electrical": [
        "electrical",
        "power",
        "electronics",
    ],

    "mechanical": [
        "mechanical",
        "machinery",
        "hvac",
    ],

    "data": [
        "data",
        "analytics",
        "sql",
    ],
}


@app.post("/api/jobs/match")
def match_jobs(
    data: SkillMatch
):

    skills_input = (
        data.skills
        .strip()
        .lower()
    )

    if not skills_input:

        raise HTTPException(
            status_code=400,
            detail="Skills are required"
        )

    if len(
        skills_input
    ) > 1000:

        raise HTTPException(
            status_code=422,
            detail=(
                "Skills must be "
                "1,000 characters or fewer."
            )
        )

    raw = []

    for part in skills_input.split(","):

        words = part.strip().split()

        for word in words:

            w = word.strip(
                '.,!?()[]{}"\''
            )

            if (
                w
                and len(w) >= 3
                and w not in FILLER_WORDS
            ):

                if w in KNOWN_SKILLS:

                    raw.append(
                        w
                    )

        if len(words) >= 2:

            for i in range(
                len(words) - 1
            ):

                phrase = (
                    f"{words[i]} "
                    f"{words[i + 1]}"
                )

                if phrase in KNOWN_SKILLS:

                    raw.append(
                        phrase
                    )

    skill_list = sorted(
        set(raw)
    )

    if not skill_list:

        for skill in KNOWN_SKILLS:

            if re.search(
                r"(?<![a-z0-9])"
                + re.escape(
                    skill
                )
                + r"(?![a-z0-9])",
                skills_input
            ):

                skill_list.append(
                    skill
                )

        skill_list.sort()

    db = SessionLocal()

    try:

        all_jobs = [
            job
            for job in db.query(
                Job
            ).all()
            if is_displayable_job(job)
        ]

        total_jobs = len(
            all_jobs
        )

        if not skill_list:

            return {
                "success": True,
                "matched_jobs": [],
                "total_matches": 0,
                "total_jobs": total_jobs,
                "match_percentage": 0,
                "extracted_skills": [],
                "message": (
                    "No valid skills detected."
                ),
                "recommendations": [
                    (
                        "Try specific skills like "
                        "'python', 'javascript', 'management'"
                    )
                ],
            }

        matched = []

        for job in all_jobs:

            text = (
                f"{job.title.lower()} "
                f"{(job.company or '').lower()} "
                f"{(job.description or '').lower()}"
            )

            match_count = 0

            matched_skills = []

            missing_skills = []

            for skill in skill_list:

                found = False

                if re.search(
                    r"(?<![a-z0-9])"
                    + re.escape(skill)
                    + r"(?![a-z0-9])",
                    text
                ):

                    match_count += 2

                    matched_skills.append(
                        skill
                    )

                    found = True

                else:

                    for key, syns in SYNONYMS.items():

                        if (
                            skill in syns
                            or skill == key
                        ):

                            for syn in syns:

                                if re.search(
                                    r"(?<![a-z0-9])"
                                    + re.escape(
                                        syn
                                    )
                                    + r"(?![a-z0-9])",
                                    text
                                ):

                                    match_count += 1

                                    matched_skills.append(
                                        f"{skill} → {syn}"
                                    )

                                    found = True

                                    break

                        if found:

                            break

                if not found:

                    missing_skills.append(
                        skill
                    )

            max_score = (
                len(skill_list)
                * 2
            )

            score = (
                (match_count / max_score)
                * 100
                if max_score
                else 0
            )

            exact = len([
                s
                for s in matched_skills
                if (
                    "→" not in s
                    and "partial" not in s
                )
            ])

            if exact > 0:

                score = min(
                    100,
                    score + 10
                )

            matched.append(
                {
                    "jobTitle": job.title,
                    "company": job.company,
                    "jobType": job.job_type,
                    "location": job.location,
                    "salary": job.salary,
                    "matchScore": round(
                        score,
                        1
                    ),
                    "matchedSkills": list(
                        set(
                            matched_skills
                        )
                    ),
                    "missingSkills": missing_skills[:3],
                    "whyMatch": (
                        "The listing mentions "
                        + ", ".join(
                            matched_skills
                        )
                        + "."
                    ),
                    "id": job.id,
                }
            )

        matched.sort(
            key=lambda x: x["matchScore"],
            reverse=True
        )

        pct = (
            round(
                (
                    len(matched)
                    / total_jobs
                )
                * 100,
                1
            )
            if total_jobs
            else 0
        )

        if len(matched) == 0:

            recs = [
                "Try different skills",
                "Make sure you have jobs in database",
                "Use specific skill names",
            ]

        elif len(matched) < 5:

            recs = [
                "Add more skills for more matches",
                "Consider broader terms",
            ]

        else:

            recs = [
                (
                    f"Great! Found "
                    f"{len(matched)} jobs "
                    f"out of "
                    f"{total_jobs}"
                ),
                (
                    "Add more skills "
                    "for better matches"
                ),
            ]

        return {
            "success": True,
            "matched_jobs": matched[:20],
            "total_matches": len(
                matched
            ),
            "total_jobs": total_jobs,
            "match_percentage": pct,
            "extracted_skills": skill_list,
            "recommendations": recs,
            "message": (
                f"Found {len(matched)} "
                f"matching jobs out of "
                f"{total_jobs} total jobs "
                f"({pct}% match rate)"
            ),
        }

    finally:

        db.close()


# ============================================================
# RUN
# ============================================================


# ============================================================
# PUBLIC JOB POSTING — Watumiaji wote wanaweza kupost
# ============================================================

class PublicJobCreate(BaseModel):
    title: str
    company: str
    job_type: Optional[str] = "onsite"
    location: Optional[str] = ""
    salary: Optional[str] = ""
    description: Optional[str] = ""
    application_url: Optional[str] = ""
    recruiter_email: Optional[str] = ""
    poster_name: Optional[str] = ""
    poster_email: Optional[str] = ""


@app.post("/api/jobs/public", status_code=201)
def create_public_job(job: PublicJobCreate):
    """
    Public job posting — no admin token required.
    Jobs are created with status 'active' so they appear immediately.
    """

    title = (job.title or "").strip()
    company = (job.company or "").strip()

    # Validation
    if not title or len(title) > 200:
        raise HTTPException(422, "Job title is required (max 200 chars).")

    if not company or len(company) > 100:
        raise HTTPException(422, "Company name is required (max 100 chars).")

    if job.job_type not in {"remote", "hybrid", "onsite", "gig"}:
        raise HTTPException(422, "Job type must be remote, hybrid, onsite, or gig.")

    if not (job.application_url or "").strip() and not (job.recruiter_email or "").strip():
        raise HTTPException(422, "Provide an application URL or recruiter email.")

    # Validate URL
    application_url = (job.application_url or "").strip()
    if application_url:
        parsed = urlparse(application_url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise HTTPException(422, "Application URL must be a valid HTTP/HTTPS link.")

    # Validate email
    recruiter_email = (job.recruiter_email or "").strip()
    if recruiter_email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", recruiter_email):
        raise HTTPException(422, "Enter a valid recruiter email address.")

    # Field length checks
    for field, value, maximum in [
        ("Description", job.description, 12000),
        ("Location", job.location, 150),
        ("Salary", job.salary, 150),
    ]:
        if value and len(value) > maximum:
            raise HTTPException(422, f"{field} must be {maximum} characters or fewer.")

    db = SessionLocal()

    try:
        new_job = Job(
            title=title,
            company=company,
            job_type=job.job_type or "onsite",
            location=(job.location or "").strip() or None,
            salary=(job.salary or "").strip() or None,
            description=(job.description or "").strip(),
            source="Public Post",
            application_url=application_url or None,
            recruiter_email=recruiter_email or None,
            job_status="active",
            posted_date=datetime.utcnow(),
            completeness_score=sum(bool(v) for v in [
                title, company, job.location, job.job_type, job.salary,
                job.description, application_url, recruiter_email,
            ]),
        )

        db.add(new_job)
        db.commit()
        db.refresh(new_job)

        print(f"✅ Public job posted: {new_job.title} at {new_job.company}")

        return {
            "success": True,
            "message": "Job posted successfully!",
            "job": {
                "id": new_job.id,
                "title": new_job.title,
                "company": new_job.company,
            },
        }

    except HTTPException:
        db.rollback()
        raise
    except Exception as error:
        db.rollback()
        print(f"Public job post failed: {error}")
        raise HTTPException(500, "The job could not be posted. Try again.")
    finally:
        db.close()


if __name__ == "__main__":

    import uvicorn

    uvicorn.run(
        app,

        host=os.environ.get(
            "HOST",
            "0.0.0.0"
        ),

        port=int(
            os.environ.get(
                "PORT",
                "5000"
            )
        ),
    )